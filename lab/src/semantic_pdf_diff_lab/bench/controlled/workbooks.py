"""The controlled corpus as Excel workbooks (.xlsx; the adapters plan, milestone 5, "Excel, in detail": "the
controlled corpus written as workbooks, scored by the same keys (the medium test: the same facts from PDF, Word,
slides and Excel)").

- **A sheet per section,** named by its heading; the document's title and any text before the first section on a
  first sheet, "Summary".
- **Paragraphs as text cells,** one a row down column A; a table's caption a cell above it; the table's cells below,
  a printed number stored as a number in its number format ("5,151" as 5151 shown "#,##0"; "96.8" as 96.8 shown
  "0.0"), so the reader shows it as Excel would.
- **A chart's data as a table** under its caption (categories by series, each series' header with its axis's unit,
  "Option 1 (MWh)", as a workbook holds the data a chart draws);
  a procedure diagram isn't carried, its caption a note saying so.
- **The key** is placed by the reader's own lines (xlsx_key).

Byte for byte the same each time (representations.FIXED_TIME).
"""
import io
import re
import zipfile

from .representations import FIXED_TIME, NOT_SHOWN

NUMBER = re.compile(r"^-?\d{1,3}(?:,\d{3})*(?:\.\d+)?$|^-?\d+(?:\.\d+)?$")
SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")

def _value(text):
    """A cell's value and number format: a printed number as a number shown as printed, anything else as text."""
    if NUMBER.match(text):
        decimals = len(text.split(".")[1]) if "." in text else 0
        grouped = "," in text
        number = float(text.replace(",", "")) if decimals else int(text.replace(",", ""))
        code = ("#,##0" if grouped else "0") + ("." + "0" * decimals if decimals else "")
        return number, code
    return text, None

def workbook(project):
    """A project's Markdown representation as a workbook's bytes."""
    import datetime
    import openpyxl
    from .corpus import charts
    from .procedures import Procedure
    from .representations import markdown
    from .word import blocks
    by_caption = {c.caption: c for c in charts(project)}
    wb = openpyxl.Workbook()
    ws, row, names = wb.active, 1, set()
    ws.title = "Summary"

    def put(r, c, text):
        value, code = _value(text)
        cell = ws.cell(r, c, value)
        if code:
            cell.number_format = code

    def sheet(title):
        nonlocal ws, row
        name = SHEET_CHARS.sub(" ", title)[:31].strip() or "Sheet"
        while name in names:
            name = name[:28] + f" {len(names)}"
        names.add(name)
        ws, row = wb.create_sheet(name), 1

    def table(rows):
        nonlocal row
        for cells in rows:
            for c, text in enumerate(cells, 1):
                if text:
                    put(row, c, text)
            row += 1
        row += 1  # a blank row after a table

    for block in blocks(markdown(project)):
        kind = block[0]
        if kind == "heading" and block[1] == 0:
            ws.cell(row, 1, block[2])
            row += 2
        elif kind == "heading":
            sheet(block[2])
        elif kind == "table":
            table(block[1])
        elif kind == "italic" and NOT_SHOWN.sub("", block[1]) in by_caption:
            figure = by_caption[NOT_SHOWN.sub("", block[1])]
            if isinstance(figure, Procedure):
                ws.cell(row, 1, NOT_SHOWN.sub("", block[1]) + " (a diagram, not shown in this workbook)")
                row += 2
                continue
            ws.cell(row, 1, figure.caption)
            row += 1
            table([["Category"] + [f"{name} ({figure.axis})" for name, _ in figure.series]] +
                  [[category] + [facts[k].value for _, facts in figure.series]
                   for k, category in enumerate(figure.categories)])
        elif kind == "bold":  # a caption: a blank row before it, the table right under it
            if row > 1 and ws.cell(row - 1, 1).value is not None:
                row += 1
            ws.cell(row, 1, block[1])
            row += 1
        else:
            ws.cell(row, 1, block[1])
            row += 1
    stamp = datetime.datetime(*FIXED_TIME)
    wb.properties.created = wb.properties.modified = stamp
    wb.properties.creator = wb.properties.lastModifiedBy = "controlled corpus"
    raw = io.BytesIO()
    wb.save(raw)
    fixed = io.BytesIO()  # the same members, each dated FIXED_TIME; openpyxl stamps the save time as modified
    when = stamp.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source, zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as out:
        for info in source.infolist():
            member = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                member = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + when, member)
            out.writestr(zipfile.ZipInfo(info.filename, FIXED_TIME), member, compress_type=zipfile.ZIP_DEFLATED)
    return fixed.getvalue()

def xlsx_key(project, data):
    """A workbook's key, placed by the reader's own lines: table rows ("xlsx-table"; a chart's facts are printed in
    its data's rows, "xlsx-chart"), cells of text ("xlsx-prose")."""
    from semantic_pdf_diff.xlsxdocs import read_xlsx
    from .corpus import charts
    from .representations import key
    doc = read_xlsx(data)
    rows = {n for b in doc.blocks if b.kind == "table" for n in b.row_lines}
    charted = rows if charts(project) else set()
    return key(project, [t for _, t in doc.lines], rows, "xlsx", None, charted)

def write(folder, project, data):
    """Write <id>.xlsx and its key into folder; returns the workbook's path."""
    import json
    from pathlib import Path
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{project.id}.xlsx"
    path.write_bytes(data)
    (folder / f"{project.id}.key.json").write_text(json.dumps(xlsx_key(project, data), indent=1, ensure_ascii=False)
                                                   + "\n", encoding="utf-8")
    return path
