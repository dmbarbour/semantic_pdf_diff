"""The controlled corpus in other representations (the controlled documents plan, decision 4; the adapters plan,
milestone 1): the same projects written as Markdown, scored by the same facts.

The owner (2026-10-02): "When we do add the alternative readers, we'll implicitly get a new form of control tests:
same facts across two or more representations, within constraints of being unable to effectively represent all
media."

- **What Markdown carries:** headings, paragraphs (the plain or trap phrasing, as the project's knob says), and
  tables as pipe tables (a schedule's sections as tables of their own, a split schedule whole).
- **What it doesn't:** charts, procedure diagrams, schematics and drawing sheets. A chart's or diagram's caption
  stays, marked as not shown; the facts only it held are listed in the key as absent, so they're neither found nor
  missed.
- **Word carries charts and diagrams as pictures** (the adapters plan, "Pictures in Word documents"): each chart
  cropped from the PDF's page as a PNG, each procedure diagram as a WMF of the same drawing (procedures.py), at its
  drawn size, its caption after it in the Caption style; its facts are placed at the picture's paragraph
  ("docx-figure") and scored.
- **The key:** the facts printed, each placed by line; every printed number logged with its role, as a PDF's are.
"""
import re

from semantic_pdf_diff.values import printed_value

from .corpus import key as pdf_key

def _row(cells):
    return "| " + " | ".join(str(c).replace("|", "\\|") for c in cells) + " |"

def _table(caption, header, rows):
    return [f"**{caption}**", "", _row(header), _row(["---"] * len(header)), *(_row(r) for r in rows), ""]

def markdown(project):
    """A project as Markdown lines, or None for one Markdown can't carry (a sheet, a schematic)."""
    if getattr(project, "sheet", None) or getattr(project, "schematic", None):
        return None
    out = [f"# {project.title}", ""]
    for heading, blocks in project.sections:
        out += [f"## {heading}", ""]
        for block in blocks:
            kind = block[0]
            if kind == "p":
                out += [block[1], ""]
            elif kind == "text":
                plain, trap = project.texts[block[1]]
                out += [(trap if project.has("traps") else plain).format(**project.values), ""]
            elif kind == "table":
                _, caption, header, rows = block
                out += _table(caption, header, rows)
            elif kind == "rooms":
                _, (caption_trap, caption_plain), header, rows = block
                trap = project.has("traps")
                shown = rows if trap else [[f"Hall C {row[0].lower()}"] + row[1:] for row in rows]
                out += _table(caption_trap if trap else caption_plain, header, shown)
            elif kind == "schedule":
                sched = block[1]
                number, _, name = sched.caption.partition(". ")
                parts = sched.sections
                for k, (title, _, columns, rows) in enumerate(parts):
                    caption = f"{number}{'abcdefgh'[k]}. {title}" if len(parts) > 1 else sched.caption
                    out += _table(caption, [title] + [c.label for c in columns], [[tag, *cells] for tag, cells in rows])
            elif kind == "chart":
                out += [f"*{block[1].caption} (a chart, not shown in this representation)*", ""]
            elif kind == "procedure":
                out += [f"*{block[1].caption} (a diagram, not shown in this representation)*", ""]
            else:
                return None  # a schematic's parts: not carried
    return out

def _numbers(line):
    """The printed numbers in a line, as words stripped of punctuation (as corpus.locate reads a page)."""
    for word in line.replace("|", " ").split():
        text = word.strip(",.;:()°*")
        if re.search(r"[0-9]", text) and printed_value(text) is not None:
            yield text

def key(project, lines, tables=None, representation="markdown", pictures=None):
    """A representation's key: the PDF key's facts that the lines print, each placed by line ("md-prose", "md-table";
    for Word, "docx-prose", "docx-table"), the others listed as absent; every printed number logged with its role.
    tables: the line numbers that are table rows (default: Markdown's, lines starting "|"). pictures: {chart number:
    the line of the picture showing it}, its facts placed there ("docx-figure")."""
    from .corpus import charts
    by_value = {}
    for f in project.facts:
        if not f.relation and f.drawn not in ("chart", "figure"):  # a chart's or diagram's facts: absent, though a
                                                                     # total may print the same number
            by_value.setdefault(printed_value(f.value), []).append(f)
        f.forms = []
    for number, line in (pictures or {}).items():
        for _, facts in charts(project)[number - 1].series:
            for f in facts:
                f.forms.append({"form": "docx-figure", "page": 1, "line": line})
    log = []
    prefix = "docx" if representation == "docx" else "md"
    for n, line in enumerate(lines, 1):
        table = n in tables if tables is not None else line.startswith("|")
        form = f"{prefix}-table" if table else f"{prefix}-prose"
        for text in _numbers(line):
            facts = by_value.get(printed_value(text), []) if re.match(r"[-+±$]?\d", text) else []
            for f in facts:
                f.forms.append({"form": form, "page": 1, "line": n})
            log.append({"text": text, "page": 1, "role": facts[0].role if facts else "structure",
                        "facts": [f.id for f in facts]})
    absent = [f.id for f in project.facts if not f.relation and not f.forms]
    out = pdf_key(project, log)
    out["facts"] = [f for f in out["facts"] if f["id"] not in absent and not f.get("relation")]
    out["absent"] = absent
    out["representation"] = representation
    return out

# Layout knobs (page furniture, columns, page breaks, rasters) mean nothing in Markdown: only the clean documents
# and the prose traps are written, with the revisions of those.
MARKDOWN_KNOBS = ("clean", "traps")

def corpus(seeds=(1,)):
    """[(project, Markdown lines)]: every corpus document Markdown carries, revisions included."""
    from .corpus import corpus as projects
    out = []
    for project in projects(seeds, knobs=True, revisions=True):
        lines = markdown(project) if project.knob in MARKDOWN_KNOBS else None
        if lines is not None:
            out.append((project, lines))
    return out

def write(project, lines, folder):
    """Write <id>.md and <id>.key.json into folder; returns the Markdown's path."""
    import json
    from pathlib import Path
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{project.id}.md"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    (folder / f"{project.id}.key.json").write_text(json.dumps(key(project, lines), indent=1, ensure_ascii=False) + "\n",
                                                   encoding="utf-8")
    return path

# --- Word (.docx) ------------------------------------------------------------------------------

FIXED_TIME = (2026, 1, 1, 0, 0, 0)  # zip members' and core properties' time, so the same project gives the same bytes

NOT_SHOWN = re.compile(r" \(a (?:chart|diagram), not shown in this representation\)$")

def chart_pictures(project, dpi=200):
    """{caption: (PNG bytes, (width, height) in points, WMF bytes or None)}: each chart as the PDF draws it, cropped
    from its page; each procedure diagram the same, with the WMF of its drawing that Word carries."""
    import pymupdf
    from .corpus import charts, render
    from .procedures import Procedure, wmf
    if not charts(project):
        return {}
    data, _ = render(project)
    doc = pymupdf.open("pdf", data)
    out = {}
    for number, chart in enumerate(charts(project), 1):
        page_no, box = project.chart_boxes[number]
        rect = pymupdf.Rect(box)
        out[chart.caption] = (doc[page_no - 1].get_pixmap(dpi=dpi, clip=rect).tobytes("png"), (rect.width, rect.height),
                              wmf(chart, rect.width) if isinstance(chart, Procedure) else None)
    return out

def docx(lines, pictures=None):
    """The Markdown lines as a Word document's bytes: headings as headings, paragraphs, pipe tables as tables, a bold
    caption line as a bold paragraph; a chart or diagram not shown, as its picture with its caption after it when
    `pictures` ({caption: (PNG, size in points, metafile or None)}) holds it: a metafile takes the PNG's place (the PNG
    sizes the picture). Byte for byte the same each time."""
    import datetime
    import io
    import zipfile
    import docx as python_docx
    document = python_docx.Document()
    props = document.core_properties
    stamp = datetime.datetime(*FIXED_TIME)
    props.created = props.modified = props.last_printed = stamp
    props.author = props.last_modified_by = "controlled corpus"
    k = 0
    while k < len(lines):
        line = lines[k]
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            document.add_heading(heading.group(2), level=len(heading.group(1)) - 1)
        elif line.startswith("|"):
            rows = []
            while k < len(lines) and lines[k].startswith("|"):
                cells = [c.strip().replace("\\|", "|") for c in lines[k].strip().strip("|").split(" | ")]
                if not all(set(c) <= {"-"} for c in cells):  # the rule under the header
                    rows.append(cells)
                k += 1
            table = document.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, text in enumerate(cells):
                    table.cell(r, c).text = text
            continue
        elif line.startswith("**") and line.endswith("**"):
            document.add_paragraph().add_run(line.strip("*")).bold = True
        elif line.startswith("*") and line.endswith("*"):
            caption = NOT_SHOWN.sub("", line.strip("*"))
            if pictures and caption in pictures:
                from docx.shared import Pt
                png, (width, _), metafile = pictures[caption]
                shape = document.add_paragraph().add_run().add_picture(io.BytesIO(png), width=Pt(width))
                if metafile:
                    _metafile(document, shape, metafile)
                document.add_paragraph(caption, style="Caption")
            else:
                document.add_paragraph().add_run(line.strip("*")).italic = True
        elif line.strip():
            document.add_paragraph(line)
        k += 1
    raw = io.BytesIO()
    document.save(raw)
    fixed = io.BytesIO()  # the same members, each dated FIXED_TIME
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source, zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as out:
        for info in source.infolist():
            out.writestr(zipfile.ZipInfo(info.filename, FIXED_TIME), source.read(info.filename),
                         compress_type=zipfile.ZIP_DEFLATED)
    return fixed.getvalue()

def _metafile(document, shape, data):
    """Point an inline picture at a WMF in place of its image, dropping the image (no other picture uses it)."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.part import Part
    from docx.oxml.ns import qn
    part = Part(document.part.package.next_partname("/word/media/image%d.wmf"), "image/x-wmf", data,
                document.part.package)
    blip = next(shape._inline.iter(qn("a:blip")))
    old = blip.get(qn("r:embed"))
    blip.set(qn("r:embed"), document.part.relate_to(part, RT.IMAGE))
    document.part.drop_rel(old)

def docx_key(project, data):
    """A Word document's key, placed by the reader's own lines (paragraphs and table rows, docxdocs.read_docx), its
    charts' facts at their pictures' paragraphs (in order: chart n, picture n)."""
    from semantic_pdf_diff.docxdocs import read_docx
    doc = read_docx(data)
    rows = {n for b in doc.blocks if b.kind == "table" for n in b.row_lines}
    pictures = {number: p.line for number, p in enumerate(doc.pictures, 1)}
    return key(project, [t for _, t in doc.lines], rows, "docx", pictures)

def write_docx(project, lines, folder):
    """Write <id>.docx and <id>.key.json into folder; returns the document's path."""
    import json
    from pathlib import Path
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    data = docx(lines, chart_pictures(project))
    path = folder / f"{project.id}.docx"
    path.write_bytes(data)
    (folder / f"{project.id}.key.json").write_text(json.dumps(docx_key(project, data), indent=1, ensure_ascii=False) + "\n",
                                                   encoding="utf-8")
    return path
