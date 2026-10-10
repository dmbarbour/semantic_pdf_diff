"""The job reading a text format (plain text, Markdown, Word, slide decks, workbooks, CSV), as extract's reads a PDF:
the document read, its sections, its text and tables made into tasks, its pictures read as images. A format's
reader, its claims' derivation and locator, and its pictures' kind are its entry in FORMATS (code review
2026-10-08, architecture 7: if-chains over formats in textdocs.py).
"""
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from functools import partial

from .schema import DerivationStep, DocxLocator, PptxLocator, Section, TextLocator, XlsxLocator, coverage_row
from .sections import SectionIndex
from .tasks import TaskCore, split_utf8
from .textdocs import parse

DERIVATION = {
    "text": [DerivationStep(step="text-file", detail="grouped paragraphs"), DerivationStep(step="model-extraction")],
    "table": [DerivationStep(step="markdown-table", detail="row with its header"), DerivationStep(step="model-extraction")],
}

def text_sections(doc, depth, pages_per_section):
    """Sections from the headings down to `depth` levels (each with its heading path), else fixed page ranges.
    Returns (sections, SectionIndex)."""
    sections, path = [], []
    for page, line, level, title in doc.headings:
        path = path[:level - 1] + [title]
        if level <= depth:
            sections.append(Section(id=f"sec{len(sections) + 1}", first_page=page, last_page=page, first_y=float(line),
                                    heading_path=list(path[:depth]), origin="headings"))
    if not sections or sections[0].first_y > 1 or sections[0].first_page > 1:  # text before the first heading
        sections.insert(0, Section(id="sec0", first_page=1, last_page=1, first_y=0.0, origin="headings" if sections
                                   else "pages"))
    if len(sections) == 1 and not doc.headings:  # no headings: page ranges
        sections = [Section(id=f"sec{k + 1}", first_page=p, last_page=min(p + pages_per_section - 1, doc.pages),
                            origin="pages") for k, p in enumerate(range(1, doc.pages + 1, pages_per_section))]
    for k, sec in enumerate(sections):  # each ends where the next begins
        nxt = sections[k + 1] if k + 1 < len(sections) else None
        sections[k] = sec.model_copy(update={"last_page": nxt.first_page if nxt else doc.pages,
                                             "last_y": nxt.first_y if nxt else None})
    return sections, SectionIndex(sections)

class TextReader:
    """The context levers' reader for a text file (extract.Context's methods, in lines rather than points)."""

    def __init__(self, doc, s):
        self.doc, self.s = doc, s
        self.by_page = defaultdict(list)
        for b in doc.blocks:
            self.by_page[b.page].append(b)
        self.tables_on = {}

    @staticmethod
    def compose(lines):
        from .context import CONTEXT_NOTE
        lines = [line for line in lines if line]
        return (CONTEXT_NOTE + "\n" + "\n".join(lines)) if lines else ""

    def for_text(self, page_no, segments, text):
        return self.compose(self.s.text_lines(self, page_no, segments, text))

    def for_table(self, page_no, bbox, flat):
        return self.compose(self.s.table_lines(self, page_no, self.table_top(page_no, bbox), flat))

    @property
    def pages(self):
        return self.doc.pages

    def page_blocks(self, page_no):
        """The page's blocks in order: [(box, text)], the text's whitespace folded, as a PDF page's blocks are."""
        return [(b.box, " ".join(b.text.split())) for b in self.by_page.get(page_no, [])]

    def table_top(self, page_no, row_box):
        for box in self.tables_on.get(page_no, ()):
            if box[1] <= row_box[1] and row_box[3] <= box[3]:
                return box
        return row_box

    def text_above(self, page_no, bbox):
        """The text of the blocks above a box on its page."""
        return " ".join(t for b, t in self.page_blocks(page_no) if b[3] <= bbox[1])

    def stem_path(self, page_no, bbox):
        """The headings a box sits under, every level ("7. Flow Control > 7.2. Stream Concurrency")."""
        path = []
        for page, line, level, title in self.doc.headings:
            if (page, line) > (page_no, bbox[1]):
                break
            path = path[:level - 1] + [title]
        return path

    def cited(self, text):
        return ""  # no glossary or figure captions read from text files yet

DOCX_DERIVATION = {
    "text": [DerivationStep(step="docx-paragraphs", detail="grouped paragraphs"), DerivationStep(step="model-extraction")],
    "table": [DerivationStep(step="docx-table", detail="row with its header"), DerivationStep(step="model-extraction")],
}

PPTX_DERIVATION = {
    "text": [DerivationStep(step="pptx-shapes", detail="grouped paragraphs"), DerivationStep(step="model-extraction")],
    "table": [DerivationStep(step="pptx-table", detail="row with its header"), DerivationStep(step="model-extraction")],
}

def chart_derivation(extension):
    return [DerivationStep(step=f"{extension.lstrip('.')}-chart", detail="a chart's cached values, a row per category"),
            DerivationStep(step="model-extraction")]

def text_locator(page, bbox, region, task):
    return TextLocator(page=page, lines=(int(bbox[1]), max(int(bbox[1]), int(bbox[3]) - 1)), region=region, task=task)

XLSX_DERIVATION = {
    "text": [DerivationStep(step="xlsx-cells", detail="grouped cells"), DerivationStep(step="model-extraction")],
    "table": [DerivationStep(step="xlsx-table", detail="row with its header"), DerivationStep(step="model-extraction")],
}

def xlsx_locator(places, page, bbox, region, task):
    """A workbook claim's locator: its lines, and the sheet and cells they were read from (their union)."""
    from openpyxl.utils.cell import get_column_letter, range_boundaries
    first, last = int(bbox[1]), max(int(bbox[1]), int(bbox[3]) - 1)
    found = [places[n] for n in range(first, last + 1) if n in places] or [("", "")]
    sheet, refs = found[0][0], [ref for s, ref in found if s == found[0][0]]
    try:
        bounds = [range_boundaries(ref) for ref in refs]
        cells = (f"{get_column_letter(min(b[0] for b in bounds))}{min(b[1] for b in bounds)}:"
                 f"{get_column_letter(max(b[2] for b in bounds))}{max(b[3] for b in bounds)}")
    except (ValueError, TypeError):  # a drawing's chart or picture, not cells
        cells = refs[0]
    return XlsxLocator(page=page, sheet=sheet, cells=cells, lines=(first, last), region=region, task=task)

CSV_DERIVATION = {
    "text": [DerivationStep(step="csv-fields", detail="grouped fields"), DerivationStep(step="model-extraction")],
    "table": [DerivationStep(step="csv-table", detail="row with its header"), DerivationStep(step="model-extraction")],
}

def csv_locator(places, page, bbox, region, task):
    """A CSV claim's locator: its lines, and the file's lines and fields they were read from (their union)."""
    from .schema import CsvLocator
    first, last = int(bbox[1]), max(int(bbox[1]), int(bbox[3]) - 1)
    found = [places[n] for n in range(first, last + 1) if n in places] or [((0, 0), (0, 0))]
    return CsvLocator(file_lines=(min(f[0][0] for f in found), max(f[0][1] for f in found)),
                      fields=(min(f[1][0] for f in found), max(f[1][1] for f in found)), lines=(first, last),
                      region=region, task=task)

def pptx_locator(page, bbox, region, task):
    return PptxLocator(page=page, lines=(int(bbox[1]), max(int(bbox[1]), int(bbox[3]) - 1)), region=region, task=task)

def docx_locator(page, bbox, region, task):
    return DocxLocator(page=page, paragraphs=(int(bbox[1]), max(int(bbox[1]), int(bbox[3]) - 1)), region=region,
                       task=task)

# Tasks fed between turns of fair share, on a page as long as a whole Word document or unpaginated text.
TASKS_PER_TURN = 20

def groups(blocks, text_bytes, section_of):
    """A page's text and code blocks grouped up to the byte budget (oversized blocks split), never across a section
    boundary: [(ids, [(box, text)])], ids like '3.0-5.0' naming the blocks, as text_groups names a PDF page's."""
    pieces = []
    for bi, block in enumerate(blocks):
        if block.kind == "table":
            continue
        for ci, chunk in enumerate(split_utf8(block.text.strip("\n"), text_bytes)):
            if chunk.strip():
                pieces.append((f"{bi}.{ci}", block.box, chunk))
    out, group, size = [], [], 0
    for piece in pieces + [None]:
        extra = len(piece[2].encode()) + 2 if piece else 0
        if group and (piece is None or size + extra > text_bytes or section_of(piece[1]) != section_of(group[-1][1])):
            ids = group[0][0] + (f"-{group[-1][0]}" if len(group) > 1 else "")
            out.append((ids, [(b, t) for _, b, t in group]))
            group, size = [], 0
        if piece:
            group.append(piece)
            size += extra
    return out

def _office(module, reader):
    """An office format's reader, imported when it reads: its libraries are the office extra's."""
    def read(raw):
        from importlib import import_module
        return getattr(import_module(f"{__package__}.{module}"), reader)(raw)
    return read

@dataclass(frozen=True)
class Format:
    """How a text format is read: its reader (bytes -> TextDocument), its claims' derivation, their locator (given the
    document read: a workbook's claims are placed by sheet and cells), and, for an office format, its pictures' kind.
    An office format's damaged file is a failed row and its pictures are read; another format's reader errors stand."""
    read: object
    derivation: dict
    locator: object
    pictures: str = ""

    @property
    def office(self):
        return bool(self.pictures)

_PLAIN, _MARKDOWN = (lambda raw: parse(raw.decode("utf-8", errors="replace"), markdown=False),
                     lambda raw: parse(raw.decode("utf-8", errors="replace"), markdown=True))
_WORKBOOK = Format(_office("xlsxdocs", "read_xlsx"), XLSX_DERIVATION, lambda doc: partial(xlsx_locator, doc.places), "xlsx")
_CSV = Format(_office("xlsxdocs", "read_csv"), CSV_DERIVATION, lambda doc: partial(csv_locator, doc.places))  # a one-sheet
FORMATS = {".txt": Format(_PLAIN, DERIVATION, lambda doc: text_locator),                     # workbook of text cells
           ".md": Format(_MARKDOWN, DERIVATION, lambda doc: text_locator),
           ".docx": Format(_office("docxdocs", "read_docx"), DOCX_DERIVATION, lambda doc: docx_locator, "docx"),
           ".pptx": Format(_office("pptxdocs", "read_pptx"), PPTX_DERIVATION, lambda doc: pptx_locator, "pptx"),
           ".xlsx": _WORKBOOK, ".xlsm": _WORKBOOK, ".csv": _CSV, ".tsv": _CSV}

def read_text(raw, extension):
    """A text format's bytes as its reader reads them (a TextDocument). Raises on a damaged office file."""
    return FORMATS[extension].read(raw)

def text_job(data, job, output, client, dispatch, progress, extension):
    """Generator doing one text file's extraction, as extract's PDF job does a PDF's: yields "page" before each page,
    then "waiting" while its requests are pending; job.state["result"] is set at the end."""
    s = client.s
    name = job.content
    form = FORMATS[extension]
    office = form.office
    core = TaskCore(job, output, client, dispatch, progress, name, text_locator, form.derivation)  # (located below)
    raw = data if isinstance(data, (bytes, bytearray)) else open(data, "rb").read()
    if office:
        try:
            doc = form.read(raw)
        except Exception as error:  # not a Word document, deck or workbook after all, or a damaged one
            core.record(coverage_row(content=job.content, task="open", status="failed",
                                     issues=[f"{name}: unreadable ({type(error).__name__}: {error})"]))
            job.state["result"] = ([], core.coverage)
            return
    else:
        doc = form.read(raw)
    core.locator = form.locator(doc)  # a workbook's claims are placed by sheet and cells, known once it's read
    if not doc.blocks:
        core.record(coverage_row(content=job.content, task="open", status="failed", issues=[f"{name}: no text"]))
        job.state["result"] = ([], core.coverage)
        return
    sections, index = text_sections(doc, s.section_depth, s.section_pages)
    core.sections, core.reader = index, TextReader(doc, s)
    if job.on_sections:
        job.on_sections(sections)
    signals, kept = {}, []  # kept: the pictures being read, their document open until their tasks are done
    for page in range(1, doc.pages + 1):
        yield "page"
        blocks = [b for b in doc.blocks if b.page == page]
        for k, (ids, segments) in enumerate(groups(blocks, s.text_bytes, lambda box, p=page: index.box(p, box).id)):
            if k and k % TASKS_PER_TURN == 0:
                yield "page"  # a long page: other sources take their turn
            core.text_task(page, segments, f"text:p{page}:{ids}")
        tables = [b for b in blocks if b.kind == "table"]
        core.reader.tables_on[page] = [b.box for b in tables]
        limit = s.rules_from()
        for ti, table in enumerate(tables):
            header, body = table.rows[0], table.rows[1:]
            width = max(len(r) for r in table.rows)

            # every loop name it uses bound now: rules answer late, after the loop has moved to another page
            # (code review 2026-10-08, B1: rows were filed under the next sheet)
            def by_itself(ri, page=page, table=table, ti=ti, header=header, body=body, width=width, checked=None):
                row = body[ri]
                if not any(c.strip() for c in row):
                    return
                line = table.row_lines[ri + 1]
                labels = table.row_headers[ri] if table.row_headers else header
                steps = None if checked is None else [core.derivation["table"][0], checked, core.derivation["table"][1]]
                core.table_task(page, (0.0, float(line), 1.0, float(line + 1)), f"table:p{page}:{ti}:{ri}", labels,
                                row, list(range(len(labels) if table.row_headers else width)),
                                derivation=chart_derivation(extension) if table.source == "chart" else steps)

            def as_table(checked=None, page=page, table=table, ti=ti, body=body, by_itself=by_itself, limit=limit):
                """The table read as one (checked: the key-value check's step, when it was asked)."""
                from . import tablegrid, tablerules
                read = by_itself if checked is None else lambda ri: by_itself(ri, checked=checked)
                if table.grid is not None and table.grid.rows and limit is not None and len(body) > limit:
                    source = DerivationStep(step=core.derivation["table"][0].step, detail="the table's cells")
                    tablerules.read(core, page, table.grid, f"rules:p{page}:{ti}", read,
                                    source if checked is None else [source, checked])
                    return
                if table.grid is not None and table.grid.rows:  # without rules: the grid's rows, and possible notes
                    for key in table.grid.keys + [r.key for r in tablegrid.possible_notes(table.grid)]:
                        read(key)
                    return
                for ri in range(len(body)):
                    read(ri)

            if not table.key_value:
                as_table()
                continue
            # a key-value list, as proposed, if the model confirms it (code review 2026-10-08, B5)
            from . import keyvalue
            g = table.grid
            boxes = [(0.0, float(n), 1.0, float(n + 1)) for n in [table.row_lines[0]] + g.lines]

            def as_pairs(checked, context, page=page, ti=ti, boxes=boxes, g=g):
                core.text_task(page, list(zip(boxes, keyvalue.pairs(g))), f"text:p{page}:kv{ti}", derivation=[
                    DerivationStep(step=core.derivation["table"][0].step, detail="a key and its value a line"),
                    checked, core.derivation["text"][-1]], context=context)
            keyvalue.read(core, page, g, f"key-value:p{page}:{ti}", boxes, as_table, as_pairs)
        seen = Counter()
        for pg, line, alt, target in doc.images:
            if pg == page:
                seen[line] += 1  # two notes at one line (a sheet's) stay two rows
                if alt in ("slide", "sheet", "workbook"):  # a deck's or workbook's note says what it is
                    issue = target[:1].upper() + target[1:]
                elif office:
                    issue = f"A picture or object in the document isn't read yet ({alt}: {target})"
                else:
                    issue = f"An image in a Markdown file isn't read: {alt or target}"
                core.record(coverage_row(content=job.content, page=page, bbox=[0.0, float(line), 1.0, float(line + 1)],
                                         task=f"image:p{page}:{line}" + (f":{seen[line]}" if seen[line] > 1 else ""),
                                         status="skipped", issues=[issue]))
        shown = [p for p in doc.pictures if p.page == page]  # a Word document is one page; a deck's, its slides
        if office and shown:  # a page's pictures follow its text
            from .pictures import Reading
            kept.append(Reading(form.pictures))
            yield from kept[-1].tasks(core, s, shown, job.content, output)
        section = index[page].id
        counts = Counter(numbers=sum(len(re.findall(r"\d", b.text)) > 0 for b in blocks), tables=len(tables),
                         characters=sum(len(b.text) for b in blocks))
        for key, value in counts.items():
            signals.setdefault(section, {}).setdefault(key, 0)
            signals[section][key] += value
    if job.on_sections:
        job.on_sections([x.model_copy(update={"signals": signals.get(x.id, {})}) for x in sections])
    while job.state["pending"]:
        yield "waiting"
    for reading in kept:
        reading.labels(core, job.content)
        reading.close()
    job.state["result"] = core.result(lambda r: (r["page"] or 0, r["task"]))
