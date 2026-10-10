"""Excel workbooks (.xlsx, .xlsm; docs/plans/multi-format-adapters-2026-09-23.md, milestone 5, "Excel, in detail"),
read through the text reader as Word documents are: a workbook becomes a TextDocument whose pages are its sheets and
whose lines are its regions' rows and cells, each line placed by its sheet and cell range.

The owner (2026-10-04): "Let's do Excel next. 99% of my spreadsheets are Excel."

- **Cells as Excel shows them:** a cell's cached value in its number format (chartxml.format_number), a date as an
  ISO date. A formula without a cached value is never evaluated: its cell reads "[formula =C6+C7, not calculated]".
- **Sheets in the workbook's order,** each a page and a section titled by its name; a hidden sheet is read, marked
  "(Hidden sheet)".
- **Regions:** Excel's defined tables first; then blocks of filled cells parted by blank rows and columns, a lone cell
  heading a wider block split off as its title, and narrow leading rows (a label and its value, "Reynolds |
  8100000", over a wider table, as the IEA workbooks stack them) split off as pairs, each a line ("Reynolds:
  8100000"). A region of one row or one column is text (titles, notes): each
  cell a line. The rest are tables on the Word reader's grid: a merged cell spans, its header rows labelled by path
  (a defined table's own count; else the first row, and the next when the first has merged cells and the next no
  number).
- **The sheet map** (TextDocument.regions): each region's sheet, range, kind and size, for the rules query.
- **An ambiguous region is skipped, never silently** (the plan's contract): a table with a column, past its first,
  that holds values under no header, as two tables pressed together do. It's recorded as "skipped: ambiguous sheet
  layout", with its range.
- **Every table is read whole,** each row a line, with its grid as rules see it (tablerules.Grid: its columns by
  letter, header labels, rows by sheet row number, the title above it). Whether its rows are read one by one or by
  rules a model writes is the text job's choice (the table_rules setting).
- **Comments** on cells are content, as Word's: after the region holding them ('Comment by Ana on B5: ...').
- **Charts** are read from their data (chartxml.py), and pictures as a Word document's are, after the sheet's
  regions.
- **Not followed or read:** external links, pivot caches, macros; a chart sheet is recorded as not read.

- **CSV files** (read_csv) are read as one-sheet workbooks of text cells, through the same regions, grids and tables,
  each line placed by the file's lines and fields.

Workbooks need openpyxl (the `office` extra); CSV files need nothing.
"""
import datetime
import io
import re
import posixpath
from dataclasses import dataclass, field

from . import chartxml, keyvalue, tablerules
from .docxdocs import Cell, _header_rows, _joined, _labels
from .textdocs import Block, Picture, TextDocument

XDR = "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
CHART = "http://schemas.openxmlformats.org/drawingml/2006/chart"
EMU_PER_POINT = 12700

@dataclass
class Region:
    """A block of a sheet: its bounds (rows and columns, from 1), its kind ("table" or "text"), for a table its header
    rows and the defined table naming it, and the lines it was read into."""
    sheet: str
    top: int
    left: int
    bottom: int
    right: int
    kind: str = "table"
    header_rows: int | None = None  # None: found by the heuristic; a defined table's own count (0: none)
    name: str = ""
    columns: tuple = ()  # a defined table's column names ("Column1"), its labels when it has no header row
    lines: list = field(default_factory=list)

    @property
    def ref(self):
        corner = f"{_letter(self.left)}{self.top}"
        return corner if (self.bottom, self.right) == (self.top, self.left) else \
            f"{corner}:{_letter(self.right)}{self.bottom}"

    @property
    def rows(self):
        return self.bottom - self.top + 1

_letter = tablerules.letter  # a column's letter, from 1

def shown(cell, formula=None):
    """A cell's value as Excel shows it: its number format applied, a date as an ISO date, an error as written; a
    formula without a cached value marked, not evaluated."""
    value = cell.value
    if value is None:
        return f"[formula {formula}, not calculated]" if isinstance(formula, str) and formula.startswith("=") else ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):  # its time shown if its format shows one, midnight too (a log's first)
        timed = re.search(r"[hs]", re.sub(r'"[^"]*"|\[[^\]]*\]', "", cell.number_format.split(";")[0].lower()))
        return value.isoformat(" ") if timed or value.time() != datetime.time() else value.date().isoformat()
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return chartxml.format_number(repr(value) if isinstance(value, float) else str(value), cell.number_format)
    return " ".join(str(value).split())

def _filled(ws, formulas):
    """{(row, column): text} of a sheet's cells that show something. Only stored cells are visited: a sheet formatted
    down to row 1,048,576 stores few, and visiting every formatted row would take minutes."""
    out = {}
    for (row, column), cell in ws._cells.items():
        text = shown(cell, formulas.get((row, column)))
        if text:
            out[(row, column)] = text
    return out

def _merges(ws):
    """{(row, column): (top, left, bottom, right)} for every cell inside a merged range."""
    out = {}
    for merged in ws.merged_cells.ranges:
        bounds = (merged.min_row, merged.min_col, merged.max_row, merged.max_col)
        for r in range(bounds[0], bounds[2] + 1):
            for c in range(bounds[1], bounds[3] + 1):
                out[(r, c)] = bounds
    return out

@dataclass
class Sheet:
    """A sheet as the reader takes it, whatever it came from (a workbook's sheet, a CSV file): its cells' text by
    (row, column) from 1, its merged cells, its comments ({(row, column): (author, text)}), Excel's defined tables
    ([(top, left, bottom, right, header rows, name)]) and whether it's hidden."""
    title: str
    cells: dict
    merges: dict = field(default_factory=dict)
    comments: dict = field(default_factory=dict)
    tables: list = field(default_factory=list)
    hidden: bool = False

def regions(sheet):
    """A sheet's regions: its defined tables; then blocks of filled (or merged-over) cells joined through shared edges
    (two tables touching only at a corner stay two), blocks whose bounds overlap merged into one (a list's sparse
    columns); a lone cell heading a wider block split off as its title."""
    cells, merges = sheet.cells, sheet.merges
    found, taken = [], set()
    for top, left, bottom, right, heads, name, columns in sheet.tables:
        # no header row is no header row (code review 2026-10-08, B9: one was taken from the data)
        found.append(Region(sheet.title, top, left, bottom, right, "table", 1 if heads is None else heads, name,
                            tuple(columns)))
        taken |= {(r, c) for r in range(top, bottom + 1) for c in range(left, right + 1)}
    occupied, boxes = (set(cells) | set(merges)) - taken, []
    while occupied:
        start = min(occupied)
        block, todo = {start}, [start]
        while todo:
            r, c = todo.pop()
            for n in ((r - 1, c), (r + 1, c), (r, c - 1), (r, c + 1)):
                if n in occupied and n not in block:
                    block.add(n)
                    todo.append(n)
        occupied -= block
        boxes.append([min(r for r, _ in block), min(c for _, c in block), max(r for r, _ in block),
                      max(c for _, c in block)])
    merged = True
    while merged:  # blocks whose bounds overlap are one region
        merged = False
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                if a[0] <= b[2] and b[0] <= a[2] and a[1] <= b[3] and b[1] <= a[3]:
                    boxes[i] = [min(a[0], b[0]), min(a[1], b[1]), max(a[2], b[2]), max(a[3], b[3])]
                    del boxes[j]
                    merged = True
                    break
            if merged:
                break
    for top, left, bottom, right in boxes:
        while top < bottom and right > left:  # a lone first-row cell over a wider block: its title
            first = [c for c in range(left, right + 1) if (top, c) in cells]
            if len(first) != 1:
                break
            found.append(Region(sheet.title, top, first[0], top, first[0], "text"))
            top += 1
            while top < bottom and not any((top, c) in cells for c in range(left, right + 1)):
                top += 1  # rows under the title merged over, empty: skipped
        pairs = top
        while pairs < bottom and right > left + 1 and all(c <= left + 1 for c in range(left, right + 1)
                                                            if (pairs, c) in cells):
            pairs += 1  # narrow leading rows (a label and its value) over a wider table: pairs
        if pairs > top and any((r, c) in cells for r in range(pairs, bottom + 1) for c in range(left + 2, right + 1)):
            found.append(Region(sheet.title, top, left, pairs - 1, left + 1, "pairs"))
            top = pairs
        held = [c for (r, c) in set(cells) | set(merges) if top <= r <= bottom and left <= c <= right]
        left, right = min(held, default=left), max(held, default=right)  # what's left, without the split rows' width
        kind = "text" if top == bottom or left == right else "table"
        found.append(Region(sheet.title, top, left, bottom, right, kind))
    return sorted(found, key=lambda g: (g.top, g.left))

def _grid(region, cells, merges):
    """[(row number, [Cell])]: a table region's rows on its columns (counted from 0). A merged cell spans its columns;
    below its first row its text repeats, continued (as a cell merged down in Word)."""
    rows = []
    for r in range(region.top, region.bottom + 1):
        out, c = [], region.left
        while c <= region.right:
            merged = merges.get((r, c))
            if merged and merged[1] != c:  # covered from the left: the spanning cell holds it
                c += 1
                continue
            right = min(merged[3], region.right) if merged else c
            if merged and merged[0] != r:
                out.append(Cell(cells.get(merged[:2], ""), c - region.left, right - region.left, continued=True))
            else:
                out.append(Cell(cells.get((r, c), ""), c - region.left, right - region.left))
            c = right + 1
        rows.append((r, out))
    return rows

def _unheaded(grid, heads):
    """Whether a column past the first holds values under no header (two tables pressed together)."""
    labelled = {k for _, row in grid[:heads] for c in row if c.text for k in range(c.first, c.last + 1)}
    used = {c.first for _, row in grid[heads:] for c in row if c.text}
    return any(k not in labelled for k in used if k > 0)

FORMULA = re.compile(rb"<(?:\w+:)?f[\s>/]")  # a cell's formula element in a worksheet part: <f>, <x:f t="shared"/>

def _sheet_parts(package):
    """{sheet name: its worksheet part's name}, from the workbook's own list."""
    book = "xl/workbook.xml"
    rels = package.rels(book)
    return {s.get("name"): rels[s.get(R + "id")][1] for s in package.xml(book).iter(S + "sheet")
            if s.get(R + "id") in rels}

class _Writer:
    """A TextDocument written sheet by sheet: lines, each placed ({line: (sheet, cell range)}), blocks, headings,
    notes of what isn't read (images), pictures and the sheet map."""

    def __init__(self):
        self.lines, self.blocks, self.headings, self.images, self.pictures = [], [], [], [], []
        self.places, self.mapped = {}, []

    def line(self, page, text, place):
        self.lines.append((page, text))
        self.places[len(self.lines)] = place
        return len(self.lines)

    def block(self, page, text, place):
        n = self.line(page, text, place)
        self.blocks.append(Block(page, n, n, text))
        return n

    def table(self, page, sheet, region, cells, merges, title):
        lines, images, blocks = self.lines, self.images, self.blocks
        grid = _grid(region, cells, merges)
        heads = min(_header_rows(grid, marked=0) if region.header_rows is None else region.header_rows, len(grid))
        row_ref = lambda r: f"{_letter(region.left)}{r}:{_letter(region.right)}{r}"
        if not region.name and _unheaded(grid, heads):
            region.kind = "ambiguous"
            where = f"{sheet}!{region.ref}" if sheet else region.ref
            images.append((page, len(lines), "sheet", f"skipped: ambiguous sheet layout ({where}): a column holds "
                                                     "values under no header, as tables pressed together do"))
            return
        labels = _labels(grid, heads)
        header_lines = [self.line(page, " | ".join(c.text for c in row), (sheet, row_ref(r))) for r, row in grid[:heads]]
        if not heads:  # a defined table without a header row: its columns' names, as its structured references use
            labels = [str(n) for n in region.columns][:len(labels)] or labels
            labels += [""] * (region.right - region.left + 1 - len(labels))
            header_lines = [self.line(page, " | ".join(labels), (sheet, row_ref(region.top)))]
        rows, row_lines, row_headers, aligned = [labels], [header_lines[0]], [], []
        for r, row in grid[heads:]:
            rows.append([c.text for c in row])
            row_headers.append([_joined(labels[c.first:c.last + 1]) for c in row])
            row_lines.append(self.line(page, " | ".join(c.text for c in row), (sheet, row_ref(r))))
            aligned.append([""] * (region.right - region.left + 1))
            for c in row:
                aligned[-1][c.first] = c.text
        first = header_lines[0]
        ruled = tablerules.grid([_letter(c) for c in range(region.left, region.right + 1)], labels, aligned,
                                [r for r, _ in grid[heads:]], row_lines[1:], title,
                                f"{sheet}!{region.ref}" if sheet else region.ref)
        # a block of two columns standing alone, perhaps a key-value list (code review 2026-10-08, B5); a defined
        # table's header is marked
        key_value = keyvalue.candidate(ruled, marked=bool(region.name) or heads != 1)
        blocks.append(Block(page, first, len(lines), "\n".join(t for _, t in lines[first - 1:]), "table", rows,
                            row_lines, row_headers, grid=ruled, key_value=key_value))
        region.lines += list(range(first, len(lines) + 1))

    def sheet(self, page, sheet, heading=True):
        """A sheet's regions read onto a page: its name a heading (unless not), "(Hidden sheet)" if hidden; pairs as
        "label: value" lines, text cell by cell, tables on their grid, comments after their region."""
        cells, merges = sheet.cells, sheet.merges
        if heading:
            n = self.block(page, sheet.title, (sheet.title, "A1"))
            self.headings.append((page, n, 1, sheet.title))
        if sheet.hidden:
            self.block(page, "(Hidden sheet)", (sheet.title, "A1"))
        above = None  # the last text cell or pairs read: a table's title when just above it
        placed = set()  # the comments read after their region
        for region in regions(sheet):
            self.mapped.append(region)
            if region.kind == "pairs":  # a label and its value: one line each
                read = []
                for r in range(region.top, region.bottom + 1):
                    shown_ = [cells[(r, c)] for c in (region.left, region.left + 1) if (r, c) in cells]
                    if shown_:
                        read.append(": ".join(shown_))
                        region.lines.append(self.block(page, read[-1], (sheet.title, f"{_letter(region.left)}"
                                                                        f"{r}:{_letter(region.left + 1)}{r}")))
                if read:  # a table's station or case, when it heads one ("Airfoil: circular; Pct Span: 0.0")
                    above = (region.bottom, region.left, "; ".join(read))
            elif region.kind == "text":
                for r in range(region.top, region.bottom + 1):
                    for c in range(region.left, region.right + 1):
                        if (r, c) in cells:
                            region.lines.append(self.block(page, cells[(r, c)], (sheet.title, f"{_letter(c)}{r}")))
                            above = (r, c, cells[(r, c)])
            else:
                title = above[2] if above and region.top - 2 <= above[0] < region.top and \
                    region.left <= above[1] <= region.right else ""
                self.table(page, sheet.title, region, cells, merges, title)
            for (r, c), (author, note) in sorted(sheet.comments.items()):
                if region.top <= r <= region.bottom and region.left <= c <= region.right:
                    self.comment(page, sheet.title, r, c, author, note)
                    placed.add((r, c))
        for (r, c), (author, note) in sorted(sheet.comments.items()):  # on a cell in no region (an empty one):
            if (r, c) not in placed:  # after the sheet's regions (code review 2026-10-08, B8: dropped)
                self.comment(page, sheet.title, r, c, author, note)

    def comment(self, page, sheet, r, c, author, note):
        self.block(page, f"Comment by {author or 'an unnamed author'} on {_letter(c)}{r}: {' '.join(note.split())}",
                   (sheet, f"{_letter(c)}{r}"))

    def document(self, pages):
        return TextDocument(max(1, pages), self.lines, self.blocks, self.headings, self.images, self.pictures,
                            self.places, self.mapped)

def read_xlsx(data):
    """A workbook's TextDocument: sheets as pages, regions' rows and cells as lines (each placed by its sheet and
    range), each sheet's name its heading; the sheet map in .regions."""
    import openpyxl
    from openpyxl.utils.cell import range_boundaries
    from .pptxdocs import Package
    raw = bytes(data)
    book = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    package = Package(raw)
    parts = _sheet_parts(package)
    # The formulas, for cells left uncalculated: loaded again only where a sheet holds a cell without a value and its
    # part a formula. Excel saves every formula's value, and the second load was half a large workbook's reading
    # (code review 2026-10-08, P1). A sheet whose part isn't found is loaded, to be safe.
    written = openpyxl.load_workbook(io.BytesIO(raw), data_only=False) if any(
        any(cell.value is None for cell in ws._cells.values())
        and (ws.title not in parts or FORMULA.search(package.read(parts[ws.title])))
        for ws in book.worksheets) else None
    out = _Writer()
    line, blocks, images, pictures = out.line, out.blocks, out.images, out.pictures

    def drawings(page, sheet):
        """A sheet's charts (read from their data) and pictures (read as a Word document's), after its regions."""
        lines = out.lines
        part = parts.get(sheet)
        drawing = package.related(part, "drawing") if part else None
        if not drawing:
            return
        for anchor in package.xml(drawing):
            for frame in anchor.iter(XDR + "graphicFrame"):
                reference = next(frame.iter("{%s}chart" % CHART), None)
                target = package.rels(drawing).get(reference.get(R + "id")) if reference is not None else None
                try:
                    found = chartxml.read(package.read(target[1]))
                    if not found.tables:
                        raise ValueError("no series")
                except Exception as error:  # a damaged or unexpected part: noted, the rest read
                    images.append((page, len(lines), "sheet", f"a chart whose data couldn't be read "
                                                             f"({type(error).__name__}: {error})"))
                    continue
                label = found.label()
                m = line(page, label, (sheet, "chart"))
                blocks.append(Block(page, m, m, label))
                for header, rows in found.tables:
                    first = line(page, " | ".join(header), (sheet, "chart"))
                    row_lines = [first] + [line(page, " | ".join(row), (sheet, "chart")) for row in rows]
                    blocks.append(Block(page, first, len(lines), "\n".join(t for _, t in lines[first - 1:]), "table",
                                        [header] + rows, row_lines, source="chart"))
            for pic in anchor.iter(XDR + "pic"):
                blip = next(pic.iter(A + "blip"), None)
                target = package.rels(drawing).get(blip.get(R + "embed")) if blip is not None else None
                if not target or target[1] not in package.names:
                    images.append((page, len(lines), "sheet", "a picture whose image isn't in the workbook"))
                    continue
                ext = next(pic.iter(A + "ext"), None)
                size = (int(ext.get("cx")) / EMU_PER_POINT, int(ext.get("cy")) / EMU_PER_POINT) \
                    if ext is not None and ext.get("cx") else None
                m = line(page, "", (sheet, "picture"))
                name = target[1]
                pictures.append(Picture(m, package.read(name), posixpath.splitext(name)[1].lower(), size, sheet,
                                        posixpath.basename(name), page))

    for page, ws in enumerate(book.worksheets, 1):
        formulas = {key: cell.value for key, cell in written[ws.title]._cells.items()
                    if isinstance(cell.value, str) and cell.value.startswith("=")} if written is not None else {}
        tables = []
        for table in ws.tables.values():  # (openpyxl's items() gives each table's range, not the table)
            left, top, right, bottom = range_boundaries(table.ref)
            tables.append((top, left, bottom, right, table.headerRowCount, table.displayName or table.name,
                           [column.name for column in table.tableColumns]))
        comments = {(cell.row, cell.column): (cell.comment.author, cell.comment.text)
                    for cell in ws._cells.values() if cell.comment}
        out.sheet(page, Sheet(ws.title, _filled(ws, formulas), _merges(ws), comments, tables,
                              ws.sheet_state != "visible"))
        drawings(page, ws.title)
    for sheet in book.chartsheets:
        images.append((1, len(out.lines), "workbook", f"a chart sheet ({sheet.title}), not read yet"))
    return out.document(len(book.worksheets))

# --- CSV: a one-sheet workbook of text cells

CSV_RECORDS = 1_000_000  # a backstop: records read from one file; past it, recorded as not read
A1 = re.compile(r"^([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?$")

def _column(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n

def _decode(raw):
    """A CSV file's text: UTF-8 (a byte-order mark dropped), else Windows-1252, as spreadsheet exports are."""
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")

def read_csv(data):
    """A CSV (or tab-separated) file's TextDocument, read as a one-sheet workbook whose cells are its fields as
    written: preamble lines are text, blocks parted by blank lines are tables (or label-and-value pairs), each read as
    a workbook's are. Each line's place is the file's lines and fields it came from ({line: ((first line, last line),
    (first field, last field))}; a quoted field may span lines). The delimiter is sniffed (comma, semicolon, tab or
    bar)."""
    import csv
    text = _decode(bytes(data))
    try:
        dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    cells, spans, before, out = {}, {}, 0, _Writer()
    for r, record in enumerate(reader, 1):
        if r > CSV_RECORDS:
            out.images.append((1, 0, "sheet", f"records past {CSV_RECORDS:,} not read (a backstop)"))
            break
        spans[r] = (before + 1, reader.line_num)
        before = reader.line_num
        for c, field_ in enumerate(record, 1):
            shown_ = " ".join(field_.split())
            if shown_:
                cells[(r, c)] = shown_
    out.sheet(1, Sheet("", cells), heading=False)
    for n, (_, ref) in out.places.items():
        m = A1.match(ref)
        first, last = (m.group(1), int(m.group(2))), (m.group(3) or m.group(1), int(m.group(4) or m.group(2)))
        out.places[n] = ((spans[first[1]][0], spans[last[1]][1]), (_column(first[0]), _column(last[0])))
    return out.document(1)
