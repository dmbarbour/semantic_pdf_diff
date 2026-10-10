"""Plain text and Markdown documents (docs/plans/multi-format-adapters-2026-09-23.md, milestone 1), and Word
documents read into the same form (docxdocs.py, milestone 2).

A text file is read as pages of lines: a form feed starts a page (as in RFCs), else the file is one page. Where a
PDF task sits on a page in a box of points, a text task sits in lines: its box is (0, first line, 1, last line + 1),
so sections, the context levers and the task core work unchanged, and its claims get a TextLocator.

- **Blocks:** paragraphs between blank lines, kept as laid out (spacing intact: text written for a monospace font,
  arrows and simple structures, reaches the model as written; the owner, 2026-10-03), Markdown pipe tables (read
  row by row, as PDF tables are), and Markdown code blocks (kept whole).
- **Sections:** Markdown headings; in plain text, numbered headings in the RFC style ("7.2.  Stream Concurrency");
  failing those, fixed page ranges.
- **Page furniture:** a line repeated near the top or bottom of three or more pages (an RFC's running header and
  footer) is left out, as page furniture isn't content.
- **Not read:** images a Markdown file links (recorded as skipped), and drawings made of characters as figures.
"""
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field

from .schema import DerivationStep, DocxLocator, PptxLocator, Section, TextLocator, XlsxLocator, coverage_row
from .sections import SectionIndex
from .tasks import TaskCore, split_utf8

MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
SETEXT = re.compile(r"^(=+|-+)\s*$")
# "7.2.  Stream Concurrency", "Appendix A.  Pseudocode", "A.1.  Details" at the left margin (an RFC's body is
# indented). A lone letter needs its dot: "A pump runs" and "I note" are sentences (code review 2026-10-08, B4).
RFC_HEADING = re.compile(r"^((?:\d+|Appendix [A-Z])(?:\.\d+)*|[A-Z](?:\.\d+)+|[A-Z](?=\.))\.?\s{1,4}(\S.{0,90})$")
FENCE = re.compile(r"^\s*(```|~~~)")
TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)*\|?\s*$")
IMAGE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")

DERIVATION = {
    "text": [DerivationStep(step="text-file", detail="grouped paragraphs"), DerivationStep(step="model-extraction")],
    "table": [DerivationStep(step="markdown-table", detail="row with its header"), DerivationStep(step="model-extraction")],
}

@dataclass
class Block:
    page: int
    first: int          # lines, counted from 1 through the whole file
    last: int
    text: str
    kind: str = "text"  # text, code or table
    rows: list = field(default_factory=list)  # a table's rows of cells, its header first
    row_lines: list = field(default_factory=list)
    row_headers: list = field(default_factory=list)  # each body row's own header labels (a Word table's merged
                                                     # cells); empty: every row is read under rows[0]
    source: str = ""    # where a table came from, if not a table: "chart" (a Word chart's data)
    grid: object = None  # a table's cells as its rules see them (tablerules.Grid): a workbook's tables, for now
    key_value: bool = False  # proposed as a key-value list (keyvalue.candidate): the model asked before it's read

    @property
    def box(self):
        return (0.0, float(self.first), 1.0, float(self.last + 1))

@dataclass
class TextDocument:
    pages: int
    lines: list         # [(page, text)] by line number - 1
    blocks: list
    headings: list      # [(page, line, level, title)]
    images: list        # [(page, line, alt, target)]: images not read (a Markdown file's links, a Word chart)
    pictures: list = field(default_factory=list)  # [Picture]: a Word document's pictures, read
    places: dict = field(default_factory=dict)    # a workbook's lines' places: {line: (sheet, cell range)}
    regions: list = field(default_factory=list)   # a workbook's sheet map: [xlsxdocs.Region]

@dataclass
class Picture:
    """A picture in a Word document: at its paragraph (line), its bytes and format, its displayed size in points
    (None: unknown), and its caption."""
    line: int
    data: bytes
    extension: str      # ".emf", ".png"
    size: tuple | None
    caption: str
    name: str           # the part's name ("image12.emf")
    page: int = 1       # its page: a slide's number in a deck (a Word document is one page)

def _cells(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]

def furniture(lines):
    """Line numbers of page furniture: a line, its digits aside, near the top or bottom of three or more pages."""
    by_page = defaultdict(list)
    for n, (page, text) in enumerate(lines, 1):
        if text.strip():
            by_page[page].append((n, text))
    edges = defaultdict(set)
    for page, rows in by_page.items():
        for n, text in rows[:2] + rows[-2:]:
            edges[re.sub(r"\d+", "#", " ".join(text.split()))].add((page, n))
    return {n for spots in edges.values() if len({p for p, _ in spots}) >= 3 for _, n in spots}

def parse(text, markdown):
    """A text file as pages, lines, blocks, headings and linked images."""
    lines, page = [], 1
    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        while "\f" in raw:
            before, _, raw = raw.partition("\f")
            if before.strip():
                lines.append((page, before.rstrip()))
            page += 1
        lines.append((page, raw.rstrip()))
    skip = furniture(lines) if page >= 3 else set()
    blocks, headings, images = [], [], []
    current, fence = [], None

    def close():
        nonlocal current
        if current:
            n0, n1 = current[0][0], current[-1][0]
            blocks.append(Block(lines[n0 - 1][0], n0, n1, "\n".join(t for _, t in current)))
        current = []
    n = 0
    while n < len(lines):
        n += 1
        pg, line = lines[n - 1]
        if n in skip:
            close()
            continue
        if markdown and fence is not None:  # inside a code block: kept whole, as written
            if FENCE.match(line):
                blocks.append(Block(pg, fence[0], n, "\n".join(fence[1]), "code"))
                fence = None
            else:
                fence[1].append(line)
            continue
        if markdown and FENCE.match(line):
            close()
            fence = (n, [])
            continue
        if not line.strip() or (current and lines[current[-1][0] - 1][0] != pg):
            close()
            if not line.strip():
                continue
        if markdown:
            for alt, target in IMAGE.findall(line):
                images.append((pg, n, alt, target))
            heading = MD_HEADING.match(line)
            nxt = lines[n][1] if n < len(lines) else ""
            if heading or (line.strip() and not current and SETEXT.match(nxt) and "|" not in line):
                close()
                level, title = (len(heading.group(1)), heading.group(2)) if heading else (1 if nxt.startswith("=") else 2, line.strip())
                headings.append((pg, n, level, title))
                blocks.append(Block(pg, n, n, line.strip()))
                if not heading:
                    n += 1  # the underline
                continue
            if "|" in line and n < len(lines) and TABLE_RULE.match(lines[n][1]):
                close()
                rows, row_lines, first = [_cells(line)], [n], n
                n += 1  # the rule
                while n < len(lines) and "|" in lines[n][1] and lines[n][1].strip():
                    n += 1
                    rows.append(_cells(lines[n - 1][1]))
                    row_lines.append(n)
                blocks.append(Block(pg, first, n, "\n".join(lines[k - 1][1] for k in range(first, n + 1)), "table",
                                    rows, row_lines))
                continue
        else:
            heading = RFC_HEADING.match(line)
            if heading and not current and not line[0].isspace():
                close()
                headings.append((pg, n, heading.group(1).count(".") + 1, " ".join(line.split())))
                blocks.append(Block(pg, n, n, " ".join(line.split())))
                continue
        current.append((n, line))
    close()
    if fence is not None:  # an unclosed code block runs to the end
        blocks.append(Block(lines[fence[0] - 1][0], fence[0], len(lines), "\n".join(fence[1]), "code"))
    return TextDocument(page, lines, blocks, headings, images)

def text_sections(doc, depth, pages_per_section):
    """Sections from the headings down to `depth` levels (each with its heading path), else fixed page ranges.
    Returns (sections, SectionIndex)."""
    sections, path = [], []
    for page, line, level, title in doc.headings:
        path = path[:level - 1] + [title]
        if level <= depth:
            sections.append(Section(id=f"sec{len(sections) + 1}", first_page=page, last_page=page, first_y=float(line),
                                    heading_path=list(path[:depth]), origin="headings"))
    if not sections or sections[0].first_y > 1 or sections[0].first_page > 1:  # text before the first heading
        sections.insert(0, Section(id="sec0", first_page=1, last_page=1, first_y=0.0, origin="headings" if sections
                                   else "pages"))
    if len(sections) == 1 and not doc.headings:  # no headings: page ranges
        sections = [Section(id=f"sec{k + 1}", first_page=p, last_page=min(p + pages_per_section - 1, doc.pages),
                            origin="pages") for k, p in enumerate(range(1, doc.pages + 1, pages_per_section))]
    for k, sec in enumerate(sections):  # each ends where the next begins
        nxt = sections[k + 1] if k + 1 < len(sections) else None
        sections[k] = sec.model_copy(update={"last_page": nxt.first_page if nxt else doc.pages,
                                             "last_y": nxt.first_y if nxt else None})
    return sections, SectionIndex(sections)

class TextReader:
    """The context levers' reader for a text file (extract.Context's methods, in lines rather than points)."""

    def __init__(self, doc, s):
        self.doc, self.s = doc, s
        self.by_page = defaultdict(list)
        for b in doc.blocks:
            self.by_page[b.page].append(b)
        self.tables_on = {}

    @staticmethod
    def compose(lines):
        from .context import CONTEXT_NOTE
        lines = [line for line in lines if line]
        return (CONTEXT_NOTE + "\n" + "\n".join(lines)) if lines else ""

    def for_text(self, page_no, segments, text):
        return self.compose(self.s.text_lines(self, page_no, segments, text))

    def for_table(self, page_no, bbox, flat):
        return self.compose(self.s.table_lines(self, page_no, self.table_top(page_no, bbox), flat))

    @property
    def pages(self):
        return self.doc.pages

    def page_blocks(self, page_no):
        """The page's blocks in order: [(box, text)], the text's whitespace folded, as a PDF page's blocks are."""
        return [(b.box, " ".join(b.text.split())) for b in self.by_page.get(page_no, [])]

    def table_top(self, page_no, row_box):
        for box in self.tables_on.get(page_no, ()):
            if box[1] <= row_box[1] and row_box[3] <= box[3]:
                return box
        return row_box

    def text_above(self, page_no, bbox):
        """The text of the blocks above a box on its page."""
        return " ".join(t for b, t in self.page_blocks(page_no) if b[3] <= bbox[1])

    def stem_path(self, page_no, bbox):
        """The headings a box sits under, every level ("7. Flow Control > 7.2. Stream Concurrency")."""
        path = []
        for page, line, level, title in self.doc.headings:
            if (page, line) > (page_no, bbox[1]):
                break
            path = path[:level - 1] + [title]
        return path

    def cited(self, text):
        return ""  # no glossary or figure captions read from text files yet

DOCX_DERIVATION = {
    "text": [DerivationStep(step="docx-paragraphs", detail="grouped paragraphs"), DerivationStep(step="model-extraction")],
    "table": [DerivationStep(step="docx-table", detail="row with its header"), DerivationStep(step="model-extraction")],
}

PPTX_DERIVATION = {
    "text": [DerivationStep(step="pptx-shapes", detail="grouped paragraphs"), DerivationStep(step="model-extraction")],
    "table": [DerivationStep(step="pptx-table", detail="row with its header"), DerivationStep(step="model-extraction")],
}

def chart_derivation(extension):
    return [DerivationStep(step=f"{extension.lstrip('.')}-chart", detail="a chart's cached values, a row per category"),
            DerivationStep(step="model-extraction")]

def text_locator(page, bbox, region, task):
    return TextLocator(page=page, lines=(int(bbox[1]), max(int(bbox[1]), int(bbox[3]) - 1)), region=region, task=task)

XLSX_DERIVATION = {
    "text": [DerivationStep(step="xlsx-cells", detail="grouped cells"), DerivationStep(step="model-extraction")],
    "table": [DerivationStep(step="xlsx-table", detail="row with its header"), DerivationStep(step="model-extraction")],
}

def xlsx_locator(places, page, bbox, region, task):
    """A workbook claim's locator: its lines, and the sheet and cells they were read from (their union)."""
    from openpyxl.utils.cell import get_column_letter, range_boundaries
    first, last = int(bbox[1]), max(int(bbox[1]), int(bbox[3]) - 1)
    found = [places[n] for n in range(first, last + 1) if n in places] or [("", "")]
    sheet, refs = found[0][0], [ref for s, ref in found if s == found[0][0]]
    try:
        bounds = [range_boundaries(ref) for ref in refs]
        cells = (f"{get_column_letter(min(b[0] for b in bounds))}{min(b[1] for b in bounds)}:"
                 f"{get_column_letter(max(b[2] for b in bounds))}{max(b[3] for b in bounds)}")
    except (ValueError, TypeError):  # a drawing's chart or picture, not cells
        cells = refs[0]
    return XlsxLocator(page=page, sheet=sheet, cells=cells, lines=(first, last), region=region, task=task)

CSV_DERIVATION = {
    "text": [DerivationStep(step="csv-fields", detail="grouped fields"), DerivationStep(step="model-extraction")],
    "table": [DerivationStep(step="csv-table", detail="row with its header"), DerivationStep(step="model-extraction")],
}

def csv_locator(places, page, bbox, region, task):
    """A CSV claim's locator: its lines, and the file's lines and fields they were read from (their union)."""
    from .schema import CsvLocator
    first, last = int(bbox[1]), max(int(bbox[1]), int(bbox[3]) - 1)
    found = [places[n] for n in range(first, last + 1) if n in places] or [((0, 0), (0, 0))]
    return CsvLocator(file_lines=(min(f[0][0] for f in found), max(f[0][1] for f in found)),
                      fields=(min(f[1][0] for f in found), max(f[1][1] for f in found)), lines=(first, last),
                      region=region, task=task)

def pptx_locator(page, bbox, region, task):
    return PptxLocator(page=page, lines=(int(bbox[1]), max(int(bbox[1]), int(bbox[3]) - 1)), region=region, task=task)

def docx_locator(page, bbox, region, task):
    return DocxLocator(page=page, paragraphs=(int(bbox[1]), max(int(bbox[1]), int(bbox[3]) - 1)), region=region,
                       task=task)

# Tasks fed between turns of fair share, on a page as long as a whole Word document or unpaginated text.
TASKS_PER_TURN = 20

def groups(blocks, text_bytes, section_of):
    """A page's text and code blocks grouped up to the byte budget (oversized blocks split), never across a section
    boundary: [(ids, [(box, text)])], ids like '3.0-5.0' naming the blocks, as text_groups names a PDF page's."""
    pieces = []
    for bi, block in enumerate(blocks):
        if block.kind == "table":
            continue
        for ci, chunk in enumerate(split_utf8(block.text.strip("\n"), text_bytes)):
            if chunk.strip():
                pieces.append((f"{bi}.{ci}", block.box, chunk))
    out, group, size = [], [], 0
    for piece in pieces + [None]:
        extra = len(piece[2].encode()) + 2 if piece else 0
        if group and (piece is None or size + extra > text_bytes or section_of(piece[1]) != section_of(group[-1][1])):
            ids = group[0][0] + (f"-{group[-1][0]}" if len(group) > 1 else "")
            out.append((ids, [(b, t) for _, b, t in group]))
            group, size = [], 0
        if piece:
            group.append(piece)
            size += extra
    return out

def read_text(raw, extension):
    """A text format's bytes as its reader reads them (a TextDocument). Raises on a damaged office file."""
    if extension == ".docx":
        from .docxdocs import read_docx
        return read_docx(raw)
    if extension == ".pptx":
        from .pptxdocs import read_pptx
        return read_pptx(raw)
    if extension in (".xlsx", ".xlsm"):
        from .xlsxdocs import read_xlsx
        return read_xlsx(raw)
    if extension in (".csv", ".tsv"):  # a CSV file, read as a one-sheet workbook of text cells
        from .xlsxdocs import read_csv
        return read_csv(raw)
    return parse(raw.decode("utf-8", errors="replace"), markdown=extension in (".md", ".markdown"))

def text_job(data, job, output, client, dispatch, progress, extension):
    """Generator doing one text file's extraction, as extract's PDF job does a PDF's: yields "page" before each page,
    then "waiting" while its requests are pending; job.state["result"] is set at the end."""
    s = client.s
    name = job.content
    word, deck, book = extension == ".docx", extension == ".pptx", extension in (".xlsx", ".xlsm")
    sheet = extension in (".csv", ".tsv")
    office = word or deck or book
    core = TaskCore(job, output, client, dispatch, progress, name,
                    docx_locator if word else pptx_locator if deck else text_locator,
                    DOCX_DERIVATION if word else PPTX_DERIVATION if deck else XLSX_DERIVATION if book else
                    CSV_DERIVATION if sheet else DERIVATION)
    raw = data if isinstance(data, (bytes, bytearray)) else open(data, "rb").read()
    if office:
        try:
            doc = read_text(raw, extension)
        except Exception as error:  # not a Word document, deck or workbook after all, or a damaged one
            core.record(coverage_row(content=job.content, task="open", status="failed",
                                     issues=[f"{name}: unreadable ({type(error).__name__}: {error})"]))
            job.state["result"] = ([], core.coverage)
            return
    else:
        doc = read_text(raw, extension)
    if sheet:
        core.locator = lambda page, bbox, region, task: csv_locator(doc.places, page, bbox, region, task)
    if book:  # a workbook's claims are placed by sheet and cells, known once it's read
        core.locator = lambda page, bbox, region, task: xlsx_locator(doc.places, page, bbox, region, task)
    if not doc.blocks:
        core.record(coverage_row(content=job.content, task="open", status="failed", issues=[f"{name}: no text"]))
        job.state["result"] = ([], core.coverage)
        return
    sections, index = text_sections(doc, s.section_depth, s.section_pages)
    core.sections, core.reader = index, TextReader(doc, s)
    if job.on_sections:
        job.on_sections(sections)
    signals, kept = {}, []  # kept: the pictures being read, their document open until their tasks are done
    for page in range(1, doc.pages + 1):
        yield "page"
        blocks = [b for b in doc.blocks if b.page == page]
        for k, (ids, segments) in enumerate(groups(blocks, s.text_bytes, lambda box, p=page: index.box(p, box).id)):
            if k and k % TASKS_PER_TURN == 0:
                yield "page"  # a long page: other sources take their turn
            core.text_task(page, segments, f"text:p{page}:{ids}")
        tables = [b for b in blocks if b.kind == "table"]
        core.reader.tables_on[page] = [b.box for b in tables]
        limit = s.rules_from()
        for ti, table in enumerate(tables):
            header, body = table.rows[0], table.rows[1:]
            width = max(len(r) for r in table.rows)

            # every loop name it uses bound now: rules answer late, after the loop has moved to another page
            # (code review 2026-10-08, B1: rows were filed under the next sheet)
            def by_itself(ri, page=page, table=table, ti=ti, header=header, body=body, width=width, checked=None):
                row = body[ri]
                if not any(c.strip() for c in row):
                    return
                line = table.row_lines[ri + 1]
                labels = table.row_headers[ri] if table.row_headers else header
                steps = None if checked is None else [core.derivation["table"][0], checked, core.derivation["table"][1]]
                core.table_task(page, (0.0, float(line), 1.0, float(line + 1)), f"table:p{page}:{ti}:{ri}", labels,
                                row, list(range(len(labels) if table.row_headers else width)),
                                derivation=chart_derivation(extension) if table.source == "chart" else steps)

            def as_table(checked=None, page=page, table=table, ti=ti, body=body, by_itself=by_itself, limit=limit):
                """The table read as one (checked: the key-value check's step, when it was asked)."""
                from . import tablerules
                read = by_itself if checked is None else lambda ri: by_itself(ri, checked=checked)
                if table.grid is not None and table.grid.rows and limit is not None and len(body) > limit:
                    source = DerivationStep(step=core.derivation["table"][0].step, detail="the table's cells")
                    tablerules.read(core, page, table.grid, f"rules:p{page}:{ti}", read,
                                    source if checked is None else [source, checked])
                    return
                if table.grid is not None and table.grid.rows:  # without rules: the grid's rows, and possible notes
                    for key in table.grid.keys + [r.key for r in tablerules.possible_notes(table.grid)]:
                        read(key)
                    return
                for ri in range(len(body)):
                    read(ri)

            if not table.key_value:
                as_table()
                continue
            # a key-value list, as proposed, if the model confirms it (code review 2026-10-08, B5)
            from . import keyvalue
            g = table.grid
            boxes = [(0.0, float(n), 1.0, float(n + 1)) for n in [table.row_lines[0]] + g.lines]

            def as_pairs(checked, context, page=page, ti=ti, boxes=boxes, g=g):
                core.text_task(page, list(zip(boxes, keyvalue.pairs(g))), f"text:p{page}:kv{ti}", derivation=[
                    DerivationStep(step=core.derivation["table"][0].step, detail="a key and its value a line"),
                    checked, core.derivation["text"][-1]], context=context)
            keyvalue.read(core, page, g, f"key-value:p{page}:{ti}", boxes, as_table, as_pairs)
        seen = Counter()
        for pg, line, alt, target in doc.images:
            if pg == page:
                seen[line] += 1  # two notes at one line (a sheet's) stay two rows
                if alt in ("slide", "sheet", "workbook"):  # a deck's or workbook's note says what it is
                    issue = target[:1].upper() + target[1:]
                elif office:
                    issue = f"A picture or object in the document isn't read yet ({alt}: {target})"
                else:
                    issue = f"An image in a Markdown file isn't read: {alt or target}"
                core.record(coverage_row(content=job.content, page=page, bbox=[0.0, float(line), 1.0, float(line + 1)],
                                         task=f"image:p{page}:{line}" + (f":{seen[line]}" if seen[line] > 1 else ""),
                                         status="skipped", issues=[issue]))
        shown = [p for p in doc.pictures if p.page == page]  # a Word document is one page; a deck's, its slides
        if office and shown:  # a page's pictures follow its text
            from .pictures import Reading
            kept.append(Reading("docx" if word else "pptx" if deck else "xlsx"))
            yield from kept[-1].tasks(core, s, shown, job.content, output)
        section = index[page].id
        counts = Counter(numbers=sum(len(re.findall(r"\d", b.text)) > 0 for b in blocks), tables=len(tables),
                         characters=sum(len(b.text) for b in blocks))
        for key, value in counts.items():
            signals.setdefault(section, {}).setdefault(key, 0)
            signals[section][key] += value
    if job.on_sections:
        job.on_sections([x.model_copy(update={"signals": signals.get(x.id, {})}) for x in sections])
    while job.state["pending"]:
        yield "waiting"
    for reading in kept:
        reading.labels(core, job.content)
        reading.close()
    job.state["result"] = core.result(lambda r: (r["page"] or 0, r["task"]))
