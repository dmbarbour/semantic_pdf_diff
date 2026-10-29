"""A report's indexes (indexes.py) and their files (indexfiles.py): docs/plans/indexes-and-glossaries-2026-10-29.md,
milestone 1. The owner, 2026-10-29: "building indexes and glossaries, e.g. of claims, diagrams, subjects", and of
the files: "An excel file is also a good representation for indices.\""""
import stubs  # noqa: F401 (a clean environment)
import csv
import io
import json
import tempfile
import unittest
from pathlib import Path

from semantic_pdf_diff import indexes, indexfiles
from semantic_pdf_diff.schema import Evidence, PdfLocator
from semantic_pdf_diff.report import unpacked
from stubs import serving
from test_analyze import asked, drawn, plant, run

def claim(cid, content, entity, attribute, value, unit, page=1):
    return Evidence(id=cid, content=content, locator=PdfLocator(page=page, bbox=(0, 0, 1, 1), region="text", task="t"),
                    entity=entity, attribute=attribute, value=value, unit=unit, kind="text", quote=f"{value} {unit}",
                    confidence=.9).model_dump()

def report(evidence, findings=(), sources=("a",), **more):
    files = [{"source": s, "path": f"{s}.pdf", "content": f"c{s}"} for s in sources]
    return {"sources": [{"name": s} for s in sources], "files": files, "evidence": list(evidence),
            "findings": list(findings), **more}

ONE = [claim("A-1", "ca", "P-101 feed pump", "rated power", "10", "kW"),
       claim("A-2", "ca", "P-101", "head", "42", "m", page=3),
       claim("A-3", "ca", "P-102 pump", "rated power", "12", "kW"),
       claim("A-4", "ca", "P-102 pump", "Rated  power", "12.5", "kW", page=9)]

class Subjects(unittest.TestCase):
    def test_a_sources_claims_by_subject_and_property(self):
        index = indexes.built(report(ONE))
        # named as its claims most often name it (the first, of equals)
        self.assertEqual([s["name"] for s in index["subjects"]], ["P-101 feed pump", "P-102 pump"])
        p101, p102 = index["subjects"]
        self.assertEqual(p101["names"], ["P-101", "P-101 feed pump"])
        self.assertEqual([p["name"] for p in p101["properties"]], ["head", "rated power"])
        self.assertEqual(p101["properties"][0]["values"], [{"source": "a", "value": "42", "unit": "m", "conditions": "",
                                                           "claim": "A-2", "where": "a.pdf, page 3"}])
        self.assertEqual([(p["name"], [v["value"] for v in p["values"]]) for p in p102["properties"]],
                         [("rated power", ["12", "12.5"])])  # its wordings one property
        self.assertEqual(indexes.properties(index)[1]["name"], "rated power")
        self.assertEqual([x["subject"] for x in indexes.properties(index)[1]["subjects"]],
                         ["P-101 feed pump", "P-102 pump"])

    def test_a_comparisons_subjects_merged_by_its_findings_else_by_name(self):
        b = [claim("B-1", "cb", "feed pump P-101", "rated power", "11", "kW"),
             claim("B-2", "cb", "fan", "air flow", "3", "m3/s"),
             claim("B-3", "cb", "P-103", "head", "50", "m")]
        a = ONE + [claim("A-5", "ca", "fan", "air flow", "2", "m3/s")]
        index = indexes.built(report(a + b, [{"a": "A-1", "b": "B-1", "relation": "different"}], ("a", "b")))
        by = {s["name"]: s for s in index["subjects"]}
        self.assertEqual(list(by), ["fan", "P-101 feed pump", "P-102 pump", "P-103"])  # A to Z, numbers as numbers
        self.assertEqual((by["P-101 feed pump"]["matched_by"], by["P-101 feed pump"]["claims"]),
                         ("findings", {"a": ["A-1", "A-2"], "b": ["B-1"]}))
        self.assertEqual((by["fan"]["matched_by"], by["fan"]["claims"]), ("name", {"a": ["A-5"], "b": ["B-2"]}))
        self.assertEqual((by["P-103"]["matched_by"], by["P-103"]["claims"]), ("", {"b": ["B-3"]}))
        page = indexfiles.markdown(index)
        self.assertIn("# Index: a ↔ b", page)
        self.assertIn("### P-101 feed pump (also feed pump P-101, P-101)\nClaims: a 2; b 1 (matched by findings)", page)
        self.assertIn("- **head**: a: 42 m (a.pdf, page 3); b: —", page)
        self.assertIn("- **rated power**: a: 10 kW (a.pdf, page 1); b: 11 kW (b.pdf, page 1)", page)

    def test_figures_and_tables(self):
        situation = {"ca": {"figures": [{"id": "f1", "page": 4, "kind": "figure", "label": "Figure 2",
                                         "caption": "Pump curve", "about": "The pumps' head against flow.",
                                         "claims": ["A-2"], "references": [{"page": 2}]}]}}
        records = [{"content": "ca", "id": "table:p9:0", "kind": "table", "label": "Table 5. Pumps", "page": 9,
                    "items": ["Tag", "Power"]},
                   {"content": "ca", "id": "note:p1:0", "kind": "note", "label": "NOTES", "page": 1}]
        evidence = [dict(ONE[3], records=["table:p9:0"])] + ONE[:3]
        index = indexes.built(report(evidence, situation=situation, records=records))
        self.assertEqual([(f["label"], f["page"], f["kind"], f["claims"], f["cited"]) for f in index["figures"]],
                         [("Figure 2", 4, "figure", ["A-2"], 1), ("Table 5. Pumps", 9, "table", ["A-4"], 0)])
        self.assertEqual(index["figures"][0]["title"], "Pump curve")
        self.assertIn('- **Figure 2 “Pump curve”** (a.pdf, page 4): The pumps\' head against flow; 1 claim; cited once',
                      indexfiles.markdown(index))

class Files(unittest.TestCase):
    def test_csv_a_row_a_value(self):
        index = indexes.built(report(ONE))
        header, rows = indexfiles.tables(index, {e["id"]: e for e in ONE})["subjects"]
        text = indexfiles.csv_text(header, rows)
        read = list(csv.reader(io.StringIO(text)))
        self.assertEqual(read[0], ["Subject", "Property", "Value", "Unit", "Conditions", "Source", "Where", "Claim"])
        self.assertEqual(read[1], ["P-101 feed pump", "head", "42", "m", "", "a", "a.pdf, page 3", "A-2"])
        self.assertEqual(len(read), 5)

    def test_the_workbook_an_index_its_subjects_collapsing(self):
        import openpyxl
        index = indexes.built(report(ONE))
        evidence = {e["id"]: e for e in ONE}
        data = indexfiles.workbook(index, evidence)
        self.assertEqual(data, indexfiles.workbook(index, evidence))  # the same bytes for the same index
        book = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(book.sheetnames, ["About", "Subjects", "Properties", "Claims", "Figures"])
        subjects = book["Subjects"]
        self.assertEqual(subjects.freeze_panes, "A2")
        self.assertEqual(subjects.auto_filter.ref, "A1:H7")  # the header, 2 subjects' rows and their 4 values
        rows = [[c.value for c in row] for row in subjects.iter_rows(min_row=2)]
        self.assertEqual(rows[0][:2], ["P-101 feed pump", "2 properties, 2 claims; also P-101"])
        self.assertEqual(rows[1], ["P-101 feed pump", "head", 42, "m", None, "a", "a.pdf, page 3", "A-2"])  # a number
        self.assertEqual([subjects.row_dimensions[n].outlineLevel for n in range(2, 8)], [0, 1, 1, 0, 1, 1])
        self.assertTrue(subjects["A2"].font.b)
        self.assertEqual(book["About"]["B1"].value, "a")
        self.assertEqual(book["Claims"].max_row, 5)

    def test_a_unit_of_dash_is_none(self):
        self.assertEqual(indexfiles.said({"value": "1.5", "unit": "-", "conditions": "at -100 deg"}),
                         "1.5 (at -100 deg)")

    def test_tags_and_long_numbers_stay_text(self):
        self.assertIn("<v>1200</v>", indexfiles.cell("A1", "1,200"))
        for text in ("007", "1234567890123456", "12 kW", "P-101"):
            self.assertIn('t="inlineStr"', indexfiles.cell("A1", text), text)

class Written(unittest.TestCase):
    def test_an_analysis_writes_its_index_beside_its_report(self):
        state = {"embedded": [], "names": 0}
        with tempfile.TemporaryDirectory() as d, serving(asked(state)) as url:
            root = Path(d)
            config = root / "settings.json"
            config.write_text(json.dumps({"search_names": False, "topic_names": False}))
            code, _, _ = run("analyze", plant(root / "plant"), "--store", root / "out", "--base-url", url, "--config",
                             config, "--no-situate")
            out = root / "out"
            data = json.loads((out / "report.json").read_text())
            page = (out / "index.md").read_text()
            files = sorted(p.name for p in (out / "index").iterdir())
            (out / "index.md").unlink()
            again, _, _ = run("report", "--store", out)  # written again with the report
            rewritten = (out / "index.md").read_text()
            book = (out / "index.xlsx").read_bytes()
            shown = unpacked((out / "report.html").read_text())
            searched = drawn((out / "report.html").as_uri() + "#subjects&q=blower")
        self.assertEqual((code, again), (0, 0))
        self.assertEqual(len(data["index"]["subjects"]), 15)
        self.assertEqual(files, ["claims.csv", "figures.csv", "properties.csv", "subjects.csv"])
        self.assertIn("# Index: plant", page)
        self.assertRegex(page, r"### blower0\n\n- \*\*discharge head\*\*: 20 m \(\S*blowers\.md, lines \d+–\d+\)")
        self.assertEqual(page, rewritten)
        self.assertEqual(book[:2], b"PK")
        self.assertIn('<button data-view="subjects">Subjects 15</button>', shown)  # the page's views
        self.assertRegex(shown, r'<div class="subject" data-text="blower0 discharge head"><h3>blower0</h3><table>'
                                r'<thead><tr><th>Property</th><th>Values</th></tr></thead><tbody><tr><td>discharge '
                                r'head</td><td><a href="#claim=[^"]+">20 m</a>')
        if searched is not None:  # a search shows the subjects holding its words
            self.assertIn("5 of 15 subjects", searched)

if __name__ == "__main__":
    unittest.main()
