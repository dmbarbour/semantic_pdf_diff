"""Continued reading (levers.ContinueReading, tasks.TaskCore.consume): an answer incomplete with its claims at the
limit is the model asking for more, and the task is asked again, told what it returned, for the rest."""
import stubs  # noqa: F401 (a clean environment)
import re
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from semantic_pdf_diff.tasks import CONTINUATION_NOTE, ExtractQuery  # noqa: E402
from semantic_pdf_diff.jobs import Job, reader_for, run_jobs
from semantic_pdf_diff.schema import Extraction  # noqa: E402
from semantic_pdf_diff.settings import Settings

TEXT = "# Plant\n\nPumps draw 110 kW, 120 kW, 130 kW, 140 kW and 150 kW.\n"

def picture_document():
    """A Word document holding one picture (a WMF) labelled with the five pumps' powers."""
    import io
    import docx
    from docx.shared import Inches
    from test_metafiles import WINDOW, wmf, wmf_text
    from test_pictures import _image_relationship, _part, _png, _qn
    document = docx.Document()
    document.add_paragraph("Pumps")
    picture = wmf(WINDOW + [wmf_text(50 + 180 * k, 100, f"{110 + 10 * k} kW") for k in range(5)])
    document.add_paragraph().add_run().add_picture(_png(), width=Inches(5))
    document.paragraphs[-1].runs[0]._r.find(".//" + _qn("a:blip")).set(
        _qn("r:embed"), document.part.relate_to(_part(document, "pumps.wmf", picture), _image_relationship()))
    document.add_paragraph("Figure 1: Pump powers", style="Caption")
    raw = io.BytesIO()
    document.save(raw)
    return raw.getvalue()

class Paging:
    """Claims every number in what it's shown, two at a time (the claim limit), leaving out those it's told it
    returned; says it's incomplete while any remain, or always (stubborn)."""
    def __init__(self, stubborn=False, **settings):
        self.s = Settings(claims_per_request=2, **settings)
        self.calls, self.cache_hits, self.usage, self.queries = 0, 0, {}, []
        self.stubborn = stubborn

    def ask(self, prompt, schema, images=(), key=None):
        query = ExtractQuery.read(prompt)
        self.queries.append(query)
        numbers = re.findall(r"\d{3}", query.data)
        rest = [n for n in numbers if f"| {n} kW" not in query.continuation]
        claims = [{"entity": "pump", "attribute": "power", "value": n, "unit": "kW", "kind": "text",
                   "quote": f"{n} kW", "confidence": 0.9} for n in rest[:2]]
        return Extraction(claims=claims, complete=len(rest) <= 2 and not self.stubborn)

def read(model, data=None, extension=".docx"):
    """(evidence, coverage) of a document read by its reader: by default, the picture document."""
    data = picture_document() if data is None else data
    with tempfile.TemporaryDirectory() as d:
        job = Job("sha256:" + "a" * 64 + extension, lambda: data, reader=reader_for(extension))
        run_jobs([[job]], Path(d), model)
    return job.state["result"]

def pictured(coverage):
    return [r for r in coverage if ":pic" in r["task"] and not r["task"].startswith("labels")]

class Continuation(unittest.TestCase):
    def test_an_answer_at_the_limit_asking_for_more_is_continued(self):
        model = Paging()
        evidence, coverage = read(model)
        self.assertEqual(sorted(e.value for e in evidence if e.locator.region == "overview"),
                         ["110", "120", "130", "140", "150"])
        asked = [q for q in model.queries if q.region == "overview"]
        self.assertEqual(len(asked), 3)  # 2 + 2 + 1 claims
        told = asked[2].continuation
        self.assertTrue(told.startswith(CONTINUATION_NOTE))
        self.assertEqual(re.findall(r"\| (\d+) kW", told), ["110", "120", "130", "140"])
        rows = pictured(coverage)
        first = rows[0]["task"]
        self.assertEqual([r["task"] for r in rows], [first, first + "-c1", first + "-c2"])
        self.assertEqual([r["status"] for r in rows], ["partial", "partial", "complete"])
        self.assertIn("Continued", rows[0]["issues"][-1])

    def test_text_tasks_are_refined_not_continued(self):
        model = Paging()
        read(model, TEXT.encode(), ".md")
        self.assertFalse(any(q.continuation for q in model.queries))

    def test_continuing_stops_at_the_limit_and_without_the_lever(self):
        model = Paging(stubborn=True, continuations=1)
        _, coverage = read(model)
        tasks = [r["task"] for r in pictured(coverage)]
        self.assertTrue(any(t.endswith("-c1") for t in tasks))
        self.assertFalse(any(t.endswith("-c2") for t in tasks))  # each task continued once at most; then its partial
        # status stands, and refinement (text split in two) may follow
        model = Paging(continuations=0)
        evidence, _ = read(model)
        self.assertFalse(any(q.continuation for q in model.queries))
        self.assertEqual(len([e for e in evidence if e.locator.region == "overview"]), 2)

    def test_an_answer_saying_the_limit_stopped_it_is_continued_too(self):
        class Saying(Paging):  # one claim short of the limit, and says why it stopped
            def ask(self, prompt, schema, images=(), key=None):
                answer = super().ask(prompt, schema, images, key)
                if answer.complete:
                    return answer
                return Extraction(claims=[c.model_dump() for c in answer.claims[:1]], complete=False,
                                  issues=["The diagram has more steps than the 2 claim limit allowed"])
        model = Saying()
        read(model)
        self.assertTrue(any(q.continuation for q in model.queries))
        clipped = Paging(stubborn=True)
        clipped.ask = lambda prompt, schema, images=(), key=None: (clipped.queries.append(ExtractQuery.read(prompt)) or
                                                                  Extraction(claims=[], complete=False,
                                                                             issues=["Text is clipped at the edge"]))
        read(clipped)
        self.assertFalse(any(q.continuation for q in clipped.queries))  # clipping isn't mended by asking again

    def test_a_continued_request_reads_back_into_its_parts(self):
        query = ExtractQuery("Do it.", "text", "1 Plant", "CONTEXT", "Pumps draw 110 kW.",
                             CONTINUATION_NOTE + "\n- pump | power | 110 kW")
        self.assertEqual(ExtractQuery.read(query.prompt()), query)
        plain = ExtractQuery("Do it.", "text", "1 Plant", "CONTEXT", "Pumps draw 110 kW.")
        self.assertEqual(ExtractQuery.read(plain.prompt()), plain)  # a first request's bytes are as before
        self.assertNotIn(CONTINUATION_NOTE, plain.prompt())

if __name__ == "__main__":
    unittest.main()
