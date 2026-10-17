"""Stance (docs/plans/claim-context-2026-10-14.md, milestone 5): a claim's attribution chain, the reported links of
who holds it, kept apart from the pipeline's own confidence."""
import stubs  # noqa: F401 (a clean environment)
import unittest

from semantic_pdf_diff.schema import Claim, Extraction, Link, claim_id

try:
    import docx  # noqa: F401 (the office extra)
    needs_docx = lambda cls: cls  # noqa: E731
except ImportError:
    needs_docx = unittest.skip("the office extra isn't installed")

TIMEOUT = dict(entity="Link failure detection", attribute="acknowledgement timeout", value="251", unit="ms",
               kind="table", quote="Proposal 1: The acknowledgement timeout is 251 ms.", confidence=0.9)

class Schema(unittest.TestCase):
    def test_a_claim_without_a_holder_keeps_its_id_and_two_holders_make_two_claims(self):
        import hashlib
        import json
        plain = Claim(**TIMEOUT)
        before = ["c", "link failure detection", "acknowledgement timeout", "251", "ms", "", False, "unknown"]
        self.assertEqual(claim_id("c", plain), "ev-" + hashlib.sha256(json.dumps(before).encode()).hexdigest()[:16])
        avelin = Claim(**TIMEOUT, stance=[Link(holder="Avelin", attitude="proposes")])
        borealis = Claim(**TIMEOUT, stance=[Link(holder="Borealis", attitude="proposes")])
        self.assertEqual(len({claim_id("c", x) for x in (plain, avelin, borealis)}), 3)
        # the attitude and a hedge describe the link; the holder names the claim
        hedged = Claim(**TIMEOUT, stance=[Link(holder="avelin", attitude="may propose", strength="may")])
        self.assertEqual(claim_id("c", hedged), claim_id("c", avelin))

    def test_a_malformed_stance_loses_the_stance_not_the_claim(self):
        answer = Extraction.model_validate({"claims": [{**TIMEOUT, "stance": "Avelin"},
                                                       {**TIMEOUT, "value": "483", "stance": [{"holder": "Borealis"},
                                                                                             {"attitude": "x"}, 7]}],
                                            "complete": True})
        self.assertEqual([[link.holder for link in c.stance] for c in answer.claims], [[], ["Borealis"]])
        self.assertTrue(answer.complete)

    def test_evidence_without_a_stance_is_written_as_before(self):
        from semantic_pdf_diff.schema import Evidence, PdfLocator
        made = lambda **k: Evidence(**TIMEOUT, **k, id="ev-1", content="c",  # noqa: E731
                                    locator=PdfLocator(page=1, bbox=(0, 0, 1, 1), region="table", task="t"))
        self.assertNotIn("stance", made().model_dump())
        self.assertEqual(made(stance=[Link(holder="Avelin")]).model_dump()["stance"],
                         [{"holder": "Avelin", "attitude": "", "strength": ""}])

class HolderColumn(unittest.TestCase):
    def test_a_tables_holder_column_names_its_rows_holder(self):
        from semantic_pdf_diff.stance import table_holder
        holder = lambda header, row: [(x.holder, x.attitude) for x in table_holder(header, row)]  # noqa: E731
        self.assertEqual(holder(["Company", "Proposals"], ["Spreadtrum [1]", "Proposal 1: ..."]),
                         [("Spreadtrum", "proposes")])
        self.assertEqual(holder(["Company", "Y/N", "Comments"], ["CMCC", "Y", ""]), [("CMCC", "")])
        self.assertEqual(holder(["Tag", "Flow (L/s)"], ["P-1", "118"]), [])  # no holder column
        self.assertEqual(holder(["Company", "Proposals"], [None, "Proposal 2: ..."]), [])  # an empty cell
        self.assertEqual(holder(["Item", "Source"], ["Timer", "ASHRAE [3, 5]"]), [("ASHRAE", "")])

    def test_a_rows_claims_take_its_holder_only_when_on(self):
        from semantic_pdf_diff.settings import Settings
        header, row = ["Company", "Proposals"], ["Avelin", "Proposal 1: The acknowledgement timeout is 251 ms."]
        self.assertEqual(Settings(table_holders=False).table_holder(header, row), ())
        self.assertEqual([x.holder for x in Settings().table_holder(header, row)], ["Avelin"])

    def test_a_holder_given_once_is_carried_down_only_when_on(self):
        # the carried_holders lever: a company named on its first row only, as merged cells print
        from semantic_pdf_diff.settings import Settings
        header = ["Company", "Proposals"]
        above = [["Avelin [1]", "Proposal 1: ..."], [None, "Proposal 2: ..."]]
        holder = lambda s: [x.holder for x in s.table_holder(header, ["", "Proposal 3: ..."], above)]  # noqa: E731
        self.assertEqual(holder(Settings(carried_holders=False)), [])
        self.assertEqual(holder(Settings()), ["Avelin"])
        self.assertEqual(holder(Settings(table_holders=False)), [])
        self.assertEqual(Settings().table_holder(header, ["", "x"], [[None, "y"]]), ())

class Views(unittest.TestCase):
    def test_a_yes_or_no_held_is_the_holders_view_of_the_proposition(self):
        from semantic_pdf_diff.stance import as_view
        view = lambda e, a, v, h: as_view(Claim(entity=e, attribute=a, value=v, kind="table", quote=v,  # noqa: E731
                                                confidence=0.9, stance=[Link(holder=h)]))
        said = lambda c: (c.entity, c.attribute, c.value, [(x.holder, x.attitude) for x in c.stance])  # noqa: E731
        self.assertEqual(said(view("CMCC", "support for Proposal 3.2.4", "Y", "CMCC")),
                         ("Proposal 3.2.4", "support", "Y", [("CMCC", "supports")]))
        self.assertEqual(said(view("Proposal 4.1.2 (Report)", "KDDI support status", "Yes", "KDDI")),
                         ("Proposal 4.1.2 (Report)", "support", "Yes", [("KDDI", "supports")]))
        self.assertEqual(said(view("HW/HiSi", "Proposal 3.1.1a (contents) support", "N", "HW/HiSi")),
                         ("Proposal 3.1.1a (contents)", "support", "N", [("HW/HiSi", "opposes")]))
        self.assertEqual(said(view("CMCC", "agreement", "Y", "CMCC")), ("CMCC", "support", "Y", [("CMCC", "supports")]))
        kept = Claim(**TIMEOUT, stance=[Link(holder="Avelin", attitude="proposes")])  # not a yes or a no
        self.assertEqual(as_view(kept), kept)
        self.assertEqual(as_view(Claim(**{**TIMEOUT, "value": "Y"})).stance, [])  # held by no one

    def test_two_companies_views_of_two_proposals_are_four_claims_only_when_on(self):
        from semantic_pdf_diff.settings import Settings
        claims = [Claim(entity=h, attribute=f"support for Proposal {p}", value="Y", kind="table", quote="Y",
                        confidence=0.9, stance=[Link(holder=h)]) for h in ("CMCC", "vivo") for p in ("3.1", "3.2")]
        on = [Settings().held(c) for c in claims]
        self.assertEqual(sorted({c.entity for c in on}), ["Proposal 3.1", "Proposal 3.2"])
        self.assertEqual(len({claim_id("c", c) for c in on}), 4)
        self.assertEqual([Settings(view_attitudes=False).held(c) for c in claims], claims)

class Asked(unittest.TestCase):
    def test_a_rows_holder_is_kept_and_the_readers_attitude_added(self):
        from semantic_pdf_diff.stance import merged
        links = lambda xs: [(x.holder, x.attitude) for x in xs]  # noqa: E731
        row, proposal = [Link(holder="Xiaomi")], [Link(holder="Xiaomi", attitude="proposes")]
        self.assertEqual(links(merged(row, [Link(holder="Xiaomi [12]", attitude="supports")])),
                         [("Xiaomi", "supports")])
        self.assertEqual(links(merged(proposal, [Link(holder="Xiaomi [12]", attitude="reports")])),
                         [("Xiaomi", "proposes")])
        self.assertEqual(links(merged(row, [])), [("Xiaomi", "")])
        self.assertEqual(links(merged((), [Link(holder="Borealis [2]", attitude="proposes")])),
                         [("Borealis", "proposes")])

    def test_the_reader_is_asked_for_a_stance_only_when_on(self):
        from semantic_pdf_diff.levers import STANCE_RULE, STANCE_SPEC
        from semantic_pdf_diff.settings import Settings
        from semantic_pdf_diff.tasks import extraction_template
        off, on = extraction_template(Settings()), extraction_template(Settings(read_holders=True))
        self.assertNotIn('"stance"', off)
        self.assertIn(STANCE_SPEC, on)
        self.assertIn(STANCE_RULE, on)
        self.assertIn(STANCE_SPEC, extraction_template(Settings(read_holders=True, split_conditions=False)))
        mine = 'List claims as JSON: {"claims": [], "complete": true, "issues": []}'
        self.assertEqual(extraction_template(Settings(extract_prompt=mine, read_holders=True)),
                         extraction_template(Settings(extract_prompt=mine)))

PROPOSALS = [["Company", "Proposals"], ["Avelin", "Proposal 1: The acknowledgement timeout is 251 ms."],
             ["Borealis", "Proposal 1: The acknowledgement timeout is 251 ms."],
             ["Corvane [3]", "Proposal 2: The report period is 754 ms."]]
VIEWS = [["Company", "Y/N", "Comments"], ["CMCC", "Y", "fine"], ["Delmar", "N", "prefer 600 ms"]]

def summary():
    """A feature lead's summary as a Word document: the companies' proposals, and their views on a proposal."""
    import io
    import docx
    d = docx.Document()
    d.add_heading("Feature lead summary", level=0)
    for heading, rows in (("Companies' proposals", PROPOSALS), ("Views on the moderator's Proposal 3.1", VIEWS)):
        d.add_heading(heading, level=1)
        table = d.add_table(rows=len(rows), cols=len(rows[0]))
        for r, cells in enumerate(rows):
            for c, text in enumerate(cells):
                table.cell(r, c).text = text
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()

class Model:
    """Reads a proposals table row by row (each 'N ms'), and a views table by rules (the company's Y or N)."""
    def __init__(self, **settings):
        from semantic_pdf_diff.settings import Settings
        self.s = Settings(**settings)
        self.calls, self.cache_hits, self.usage, self.queries = 0, 0, {}, []

    def ask(self, prompt, schema, images=(), key=None):
        import re
        from semantic_pdf_diff.keyvalue import Check
        from semantic_pdf_diff.tablerules import Review, Rules
        from semantic_pdf_diff.tasks import ExtractQuery
        if schema is Check:  # two columns: a table, not a key-value list
            return Check(reading="table")
        if schema is Review:  # the rules shown again with what they gave: kept
            return Review.model_validate({"verdict": "keep"})
        if schema is Rules:
            if "Y/N" not in prompt:
                return Rules(reading="rows")
            return Rules.model_validate({
                "reading": "rules", "claims": [{"columns": "", "entity": "{A}", "attribute": "agreement",
                                                "value": "{B}"}],
                "examples": [{"row": "1", "claims": [{"entity": "CMCC", "attribute": "agreement", "value": "Y"}]}]})
        query = ExtractQuery.read(prompt)
        claims = [{"entity": "Link failure detection", "attribute": "timeout", "value": v, "unit": "ms",
                   "kind": "table", "quote": v, "confidence": 0.9} for v in re.findall(r"(\d+) ms", query.data)]
        return Extraction(claims=claims, complete=True)

class RowsModel(Model):
    """Reads every table row by row: a views row gives only its comment's number, never its Y or N."""
    def ask(self, prompt, schema, images=(), key=None):
        from semantic_pdf_diff.tablerules import Rules
        return Rules(reading="rows") if schema is Rules else super().ask(prompt, schema, images, key)

class CommentRules(Model):
    """Reads the views table by a rule that takes the comment only, never the Y or N."""
    def ask(self, prompt, schema, images=(), key=None):
        from semantic_pdf_diff.tablerules import Rules
        if schema is Rules and "Y/N" in prompt:
            return Rules.model_validate({
                "reading": "rules", "claims": [{"columns": "", "entity": "{A}", "attribute": "comment",
                                                "value": "{C}"}],
                "examples": [{"row": "1", "claims": [{"entity": "CMCC", "attribute": "comment", "value": "fine"}]}]})
        return super().ask(prompt, schema, images, key)

@needs_docx
class Read(unittest.TestCase):
    def read(self, model=Model, **settings):
        import test_textdocs
        model_class, test_textdocs.Model = test_textdocs.Model, model
        try:
            evidence, _, _, _ = test_textdocs.extract("summary.docx", summary(), **settings)
        finally:
            test_textdocs.Model = model_class
        return [(e.value, [(x.holder, x.attitude) for x in e.stance]) for e in evidence if e.kind == "table"]

    def test_a_tables_rows_claims_are_held_by_its_companies_only_when_on(self):
        on = self.read()
        # row by row: two companies' proposals of one value are two claims; a reference mark left out
        self.assertEqual(sorted(x for x in on if x[0] != "Y" and x[0] != "N"),
                         [("251", [("Avelin", "proposes")]), ("251", [("Borealis", "proposes")]),
                          ("754", [("Corvane", "proposes")])])
        # by rules: the views table's rows, each its company's view (view_attitudes: the attitude from the value)
        self.assertIn(("Y", [("CMCC", "supports")]), on)
        self.assertIn(("N", [("Delmar", "opposes")]), on)
        self.assertIn(("Y", [("CMCC", "")]), self.read(view_attitudes=False))
        off = self.read(table_holders=False)
        self.assertTrue(all(stance == [] for _, stance in off))
        self.assertEqual(sorted(v for v, _ in off if v not in ("Y", "N")), ["251", "754"])  # one 251: one claim


    def test_a_rows_yes_or_no_is_its_holders_view_where_the_reader_gives_none_only_when_on(self):
        # the view_rows lever: read row by row, the reader gives the comment's number only; read by rules, the
        # comment only. Either way the row's Y or N is its company's view of the item named above the table
        views = [("N", [("Delmar", "opposes")]), ("Y", [("CMCC", "supports")])]
        for model in (RowsModel, CommentRules):
            with self.subTest(model=model.__name__):
                on = self.read(model)
                self.assertEqual(sorted(x for x in on if x[0] in ("Y", "N")), views)
                self.assertEqual([x for x in self.read(model, view_rows=False) if x[0] in ("Y", "N")], [])
        # where the reader gives the yes or no itself (the rule reads it as "agreement"), nothing is added
        self.assertEqual(sorted(x for x in self.read() if x[0] in ("Y", "N")), views)

    def test_a_view_needs_a_holder_a_yes_or_no_and_an_item_named_above(self):
        from semantic_pdf_diff.stance import view_of
        header, above = ["Company", "Y/\nN", "Comments"], "Above the table: ...Proposal 2.2: x. Table 2. Views"
        self.assertEqual(view_of(header, ["vivo", "Yes", ""], above).entity, "Proposal 2.2")
        self.assertIsNone(view_of(header, ["vivo", "Fine", ""], above))  # not a yes or a no
        self.assertIsNone(view_of(header, ["vivo", "Y", ""], "Above the table: ...Table 2. Views"))  # no item
        self.assertIsNone(view_of(["Option", "Y/N"], ["A", "Y"], above))  # no holder column


class Hedges(unittest.TestCase):
    def test_a_hedge_in_conditions_is_moved_to_uncertainty_only_when_on(self):
        from semantic_pdf_diff.settings import Settings
        from semantic_pdf_diff.stance import hedges_apart
        hedged = Claim(**{**TIMEOUT, "conditions": "link idle; status: Working assumption"})
        moved = hedges_apart(hedged)
        self.assertEqual((moved.conditions, moved.uncertainty), ("link idle", "Working assumption"))
        self.assertEqual(Settings().stated(hedged), moved)
        self.assertEqual(Settings(hedges_apart=False).stated(hedged), hedged)
        self.assertEqual(claim_id("c", moved), claim_id("c", Claim(**{**TIMEOUT, "conditions": "link idle"})))
        plain = Claim(**{**TIMEOUT, "conditions": "at peak flow"})
        self.assertEqual(hedges_apart(plain), plain)
        both = hedges_apart(Claim(**{**TIMEOUT, "conditions": "Preliminary estimates", "uncertainty": "± 5 ms"}))
        self.assertEqual((both.conditions, both.uncertainty), ("", "± 5 ms; Preliminary estimates"))


class Voice(unittest.TestCase):
    def test_the_documents_own_voice_is_no_holder(self):
        from semantic_pdf_diff.stance import merged
        read = lambda h: [Link(holder=h, attitude="proposes")]  # noqa: E731
        for voice in ("moderator", "The Moderator", "Moderator (Samsung)", "design team", "the authors", "FL", "we"):
            with self.subTest(voice):
                self.assertEqual(merged((), read(voice)), [])
        self.assertEqual([x.holder for x in merged((), read("Avelin [1]"))], ["Avelin"])
        self.assertEqual([x.holder for x in merged((), read("Moderna"))], ["Moderna"])  # a company, not a moderator


class Examples(unittest.TestCase):
    def test_the_hedge_rule_and_the_worked_examples_only_when_on(self):
        from semantic_pdf_diff.levers import HEDGE_RULE, worked_examples
        from semantic_pdf_diff.settings import Settings
        from semantic_pdf_diff.tasks import extraction_template
        plain = extraction_template(Settings())
        self.assertNotIn(HEDGE_RULE, plain)
        self.assertNotIn("WORKED EXAMPLES", plain)
        self.assertIn(HEDGE_RULE, extraction_template(Settings(hedge_rule=True)))
        both = extraction_template(Settings(hedge_rule=True, worked_examples=True))
        self.assertTrue(both.startswith(extraction_template(Settings(hedge_rule=True))))  # after: a shared prefix
        self.assertTrue(both.endswith(worked_examples()))
        # the examples' fields are the template's in force: no frame without split_conditions; a stance with holders
        self.assertNotIn('"frame"', worked_examples(frame=False))
        self.assertIn('"conditions": "on level 2; to face of finish"', worked_examples(frame=False))
        self.assertIn('"holder": "Halvern"', worked_examples(holders=True))
        self.assertNotIn("Halvern", worked_examples())

    def test_the_hedge_rule_and_the_stance_rule_alike_in_either_lever_order(self):
        from semantic_pdf_diff import levers as L
        from semantic_pdf_diff.settings import settings_class
        from semantic_pdf_diff.tasks import extraction_template
        on = {"read_holders": True, "hedge_rule": True}
        forwards = extraction_template(settings_class(L.DEFAULT_LEVERS)(**on))
        self.assertEqual(extraction_template(settings_class(tuple(reversed(L.DEFAULT_LEVERS)))(**on)), forwards)
        self.assertIn(L.NEGATION_RULE + L.HEDGE_RULE + L.STANCE_RULE, forwards)

if __name__ == "__main__":
    unittest.main()
