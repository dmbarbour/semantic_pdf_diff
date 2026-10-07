"""Slide decks (pptxdocs.py; the adapters plan, milestone 3): read as Word documents are, a slide a page, its shapes
in reading order, tables on their grid, charts from their data, pictures with their slide, speaker notes after."""
import stubs  # noqa: F401 (a clean environment)
import io
import unittest

try:
    import docx  # noqa: F401 (the office extra: the deck reader shares the Word reader's parts)
    from PIL import Image
except ImportError:
    docx = None

A = "http://schemas.openxmlformats.org/drawingml/2006/main"

def frame(x, y, cx, cy):
    return f'<a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'

def box(x, y, text, name="Box"):
    """A text box at (x, y), its paragraph."""
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="9" name="{name}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>'
            f'{frame(x, y, 2000000, 400000)}</p:spPr><p:txBody><a:bodyPr/><a:p><a:r><a:t>{text}</a:t></a:r></a:p>'
            '</p:txBody></p:sp>')

def deck():
    """A deck: a titled slide of shapes out of reading order, a numbered list, a group and a connector; a slide with a
    merged table and speaker notes; an untitled slide with a chart and a picture; a hidden slide whose title takes
    its place from its layout."""
    from semantic_pdf_diff_lab.bench.controlled.corpus import Chart, Fact
    from semantic_pdf_diff_lab.bench.controlled.slides import Deck
    from semantic_pdf_diff_lab.bench.controlled.word import chart_xml
    d = Deck()
    one = d.slide("1 Design Basis")
    one.shapes.append(box(6000000, 2000000, "Peak flow is 42.2 MGD."))      # right: read after the left one
    one.shapes.append(box(500000, 2000000, "Design flow is 25.9 MGD."))
    one.shapes.append(box(500000, 4000000, "Alum is dosed at 38.8 mg/L."))  # lower: read last
    one.text(["Screen the raw water", "Dose alum at 38.8 mg/L"], numbered=True)
    one.shapes.append(  # a group whose children's own frame is scaled onto the slide
        '<p:grpSp><p:nvGrpSpPr><p:cNvPr id="20" name="Group"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr>'
        f'<a:xfrm><a:off x="500000" y="5000000"/><a:ext cx="1000000" cy="500000"/><a:chOff x="0" y="0"/>'
        f'<a:chExt cx="2000000" cy="1000000"/></a:xfrm></p:grpSpPr>{box(0, 0, "Basin A")}{box(1000000, 0, "Basin B")}'
        '</p:grpSp>')
    one.shapes.append('<p:cxnSp><p:nvCxnSpPr><p:cNvPr id="30" name="Arrow"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>'
                      f'<p:spPr>{frame(800000, 5200000, 900000, 0)}</p:spPr></p:cxnSp>')
    two = d.slide("2 Chillers")
    two.table([["Unit", "Capacity (tons)", ""], ["", "Summer", "Winter"], ["Chiller CH-1", "450", "300"],
               ["", "425", ""]], merges=[(0, 0, 1, 0), (0, 1, 0, 2), (2, 0, 3, 0), (3, 1, 3, 2)])
    two.notes.append("Chiller CH-1 was retested in June.")
    three = d.slide()
    fact = lambda v: Fact("x", "x", (), "x", (), v)
    three.chart(chart_xml(Chart("Figure 1", ["May", "June"], [("Option 1", [fact("1,750"), fact("300")])], "MWh",
                                50, 2000)))
    png = io.BytesIO()
    Image.new("RGB", (40, 20), (200, 30, 30)).save(png, format="PNG")
    three.picture(png.getvalue(), "png", (200.0, 100.0))
    four = d.slide(hidden=True)  # its title has no frame of its own: the layout's
    four.shapes.append('<p:sp><p:nvSpPr><p:cNvPr id="2" name="Title"/><p:cNvSpPr/><p:nvPr><p:ph type="title"/>'
                       '</p:nvPr></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:p><a:r><a:t>Backup</a:t></a:r></a:p>'
                       '</p:txBody></p:sp>')
    four.shapes.append(box(500000, 300000, "Above the title, but read after it."))
    layout = (f'<p:sldLayout xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="{A}">'
              '<p:cSld><p:spTree><p:sp><p:nvSpPr><p:cNvPr id="2" name="Title"/><p:cNvSpPr/><p:nvPr><p:ph type="title"/>'
              f'</p:nvPr></p:nvSpPr><p:spPr>{frame(500000, 1000000, 8000000, 600000)}</p:spPr></p:sp></p:spTree>'
              '</p:cSld></p:sldLayout>')
    d.parts["ppt/slideLayouts/slideLayout1.xml"] = layout.encode()
    four.rels.append(("rId9", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout",
                      "../slideLayouts/slideLayout1.xml"))
    return d.save()

@unittest.skipIf(docx is None, "python-docx and Pillow aren't installed (the office extra)")
class Reading(unittest.TestCase):
    def test_slides_are_pages_their_shapes_in_reading_order(self):
        from semantic_pdf_diff.pptxdocs import read_pptx
        doc = read_pptx(deck())
        self.assertEqual(doc.pages, 4)
        self.assertEqual([(page, title) for page, _, _, title in doc.headings],
                         [(1, "1 Design Basis"), (2, "2 Chillers"), (3, "Slide 3"), (4, "Backup")])
        one = [t for page, t in doc.lines if page == 1]
        self.assertEqual(one, ["1 Design Basis", "1. Screen the raw water", "2. Dose alum at 38.8 mg/L",  # the top
                               "Design flow is 25.9 MGD.", "Peak flow is 42.2 MGD.",  # one band, left to right
                               "Alum is dosed at 38.8 mg/L.", "Basin A", "Basin B"])
        self.assertEqual([(page, what) for page, _, _, what in doc.images],
                         [(1, "a diagram of shapes joined by 1 connector(s): its text read, not its arrangement")])
        four = [t for page, t in doc.lines if page == 4]
        self.assertEqual(four, ["Backup", "(Hidden slide)", "Above the title, but read after it."])

    def test_tables_charts_pictures_and_notes(self):
        from semantic_pdf_diff.pptxdocs import read_pptx
        doc = read_pptx(deck())
        chillers, chart = [b for b in doc.blocks if b.kind == "table"]
        self.assertEqual(chillers.rows, [["Unit", "Capacity (tons) > Summer", "Capacity (tons) > Winter"],
                                         ["Chiller CH-1", "450", "300"], ["Chiller CH-1", "425"]])
        self.assertEqual(chillers.row_headers[1], ["Unit", "Capacity (tons) > Summer / Winter"])
        self.assertEqual((chillers.grid.place, chillers.grid.title), ("slide 2, table 1", "2 Chillers"))
        self.assertIsNone(chart.grid)  # a chart's data is read row by row
        self.assertIn("Speaker notes: Chiller CH-1 was retested in June.", [t for page, t in doc.lines if page == 2])
        self.assertEqual((chart.page, chart.source, chart.rows), (3, "chart", [["Category", "Option 1"],
                                                                              ["May", "1,750"], ["June", "300"]]))
        picture, = doc.pictures
        self.assertEqual((picture.page, picture.extension, picture.size), (3, ".png", (200.0, 100.0)))

    def test_a_deck_is_extracted_with_slide_locators(self):
        import sys
        from pathlib import Path
        sys.path.insert(0, str(Path(__file__).parent))
        from test_textdocs import extract
        from semantic_pdf_diff.models import PptxLocator
        evidence, coverage, sections, _ = extract("deck.pptx", deck())
        values = {e.value: e for e in evidence}
        self.assertIsInstance(values["25.9"].locator, PptxLocator)
        self.assertEqual(values["25.9"].locator.page, 1)
        self.assertEqual(values["1,750"].locator.page, 3)
        self.assertEqual([d.step for d in values["1,750"].derivation][:1], ["pptx-chart"])
        self.assertEqual([s.heading_path for s in sections][:2], [["1 Design Basis"], ["2 Chillers"]])

if __name__ == "__main__":
    unittest.main()
