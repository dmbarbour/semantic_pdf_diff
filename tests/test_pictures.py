"""A Word document's pictures (pictures.py; the adapters plan, "Pictures in Word documents"): each read as a PDF's
figure is, its caption as source text and its own text as a text layer, its claims located at its paragraph."""
import stubs  # noqa: F401 (a clean environment)
import io
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

try:
    import docx
    import pyclipper  # noqa: F401
    from PIL import Image
except ImportError:  # the office extra
    docx = None

def build():
    """A Word document: a paragraph of text, then three pictures, each captioned: a WMF schematic labelled "Flow rate
    75 gpm", a PNG, and a WMF that's no metafile at all."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml.ns import qn
    from docx.shared import Inches
    from test_metafiles import WINDOW, wmf, wmf_text
    document = docx.Document()
    document.add_heading("1 Process", level=1)
    document.add_paragraph("Pump P-1 draws 120 kW.")

    def picture(data, name, caption):
        png = io.BytesIO()
        Image.new("RGB", (40, 20), (200, 30, 30)).save(png, format="PNG")
        shape = document.add_paragraph().add_run().add_picture(io.BytesIO(png.getvalue()), width=Inches(2))
        if data is not None:  # a metafile in place of the PNG
            part = Part(PackURI(f"/word/media/{name}"), "image/x-wmf", data, document.part.package)
            next(shape._inline.iter(qn("a:blip"))).set(qn("r:embed"), document.part.relate_to(part, RT.IMAGE))
        document.add_paragraph(caption, style="Caption")
    picture(wmf(WINDOW + [wmf_text(100, 100, "Flow rate 75 gpm")]), "schematic.wmf", "Figure 1: Flow schematic")
    picture(None, "logo.png", "Figure 2: Logo")
    picture(b"not a metafile", "broken.wmf", "Figure 3: Broken")
    raw = io.BytesIO()
    document.save(raw)
    return raw.getvalue()

def _png():
    png = io.BytesIO()
    Image.new("RGB", (40, 20), (200, 30, 30)).save(png, format="PNG")
    return io.BytesIO(png.getvalue())

def _qn(tag):
    from docx.oxml.ns import qn
    return qn(tag)

def _part(document, name, data):
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    return Part(PackURI(f"/word/media/{name}"), "image/x-wmf", data, document.part.package)

def _image_relationship():
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    return RT.IMAGE

@unittest.skipIf(docx is None, "needs the office extra (python-docx, Pillow, pyclipper)")
class Pictures(unittest.TestCase):
    def test_pictures_are_read_with_their_captions_at_their_paragraphs(self):
        from semantic_pdf_diff.docxdocs import read_docx
        from semantic_pdf_diff.models import DocxLocator
        from test_textdocs import extract
        data = build()
        doc = read_docx(data)
        self.assertEqual([(p.extension, p.caption) for p in doc.pictures],
                         [(".wmf", "Figure 1: Flow schematic"), (".png", "Figure 2: Logo"),
                          (".wmf", "Figure 3: Broken")])
        schematic = doc.pictures[0]
        self.assertEqual([round(v) for v in schematic.size], [144, 72])  # 2 inches wide, as displayed
        evidence, coverage, sections, model = extract("process.docx", data)
        # the schematic's label, read from its text layer, located at its paragraph with the crop it was read from
        flow = next(e for e in evidence if e.value == "75")
        self.assertIsInstance(flow.locator, DocxLocator)
        self.assertEqual((flow.locator.paragraphs, flow.locator.region), ((schematic.line, schematic.line), "overview"))
        self.assertTrue(flow.image and flow.image.startswith("assets/"))
        self.assertEqual(flow.section, sections[-1].id)
        asked = next(q for q in model.queries if q.region == "overview" and "Flow schematic" in q.data)
        self.assertIn("Caption: Figure 1: Flow schematic", asked.data)
        self.assertIn("Flow rate 75 gpm", asked.data)  # the picture's own text, as a PDF figure's text layer
        statuses = {r["task"]: (r["status"], r["issues"]) for r in coverage if "pic" in r["task"]}
        self.assertEqual(statuses[f"overview:p1:pic{schematic.line}"][0], "complete")
        self.assertEqual(statuses[f"overview:p1:pic{doc.pictures[1].line}"][0], "complete")
        broken, issues = statuses[f"picture:p1:{doc.pictures[2].line}"]
        self.assertEqual(broken, "skipped")
        self.assertIn("broken.wmf: not drawn", issues[0])
        # the label check: the model claimed the number, but none of the picture's words ("flow", "rate", "gpm")
        status, issues = statuses[f"labels:p1:pic{schematic.line}"]
        self.assertEqual(status, "complete")  # a quality measure, not a failure
        self.assertEqual(issues, ["Label coverage 0%: 3 of the picture's 3 words in no claim (flow, gpm, rate)"])
        self.assertNotIn(f"labels:p1:pic{doc.pictures[1].line}", statuses)  # the PNG has no text layer

    def test_large_pictures_are_tiled_and_their_tiles_refined_after_every_picture_is_drawn(self):
        # through the CLI and a model answering out of order, some tiles partial: they're refined later, from pages
        # that later pictures were added after (once, adding a page invalidated them)
        import contextlib, json, tempfile
        from docx.shared import Inches
        from test_concurrency import jittery_model
        from test_metafiles import WINDOW, wmf, wmf_text
        from semantic_pdf_diff import cli
        document = docx.Document()
        document.add_paragraph("Pump P-1 draws 120 kW.")
        for n in range(3):  # wide pictures, read in tiles
            picture = wmf(WINDOW + [wmf_text(100 + 200 * k, 100 + 100 * (k % 3), f"{n}{k} kW") for k in range(4)])
            path = Path(tempfile.mkdtemp()) / f"wide{n}.wmf"
            path.write_bytes(picture)
            document.add_paragraph().add_run().add_picture(_png(), width=Inches(9))
            document.paragraphs[-1].runs[0]._r.find(".//" + _qn("a:blip")).set(_qn("r:embed"), document.part.relate_to(
                _part(document, f"wide{n}.wmf", picture), _image_relationship()))
            document.add_paragraph(f"Figure {n + 1}: Wide schematic", style="Caption")
        raw = io.BytesIO()
        document.save(raw)
        with tempfile.TemporaryDirectory() as d, jittery_model() as (url, state):
            root = Path(d)
            (root / "a.docx").write_bytes(raw.getvalue())
            args = [str(root / "a.docx"), str(root / "a.docx"), "--out", str(root / "out"), "--base-url", url, "-q"]
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                self.assertIn(cli.main(args), (0, 2))
            report = json.loads((root / "out" / "report.json").read_text())
        tasks = [r["task"] for r in report["coverage"]]
        self.assertTrue(any(t.startswith("tile:p1:pic") for t in tasks))
        self.assertTrue(any(t.startswith("tile:p1:pic") and "-r" in t for t in tasks))  # refined tiles

    def test_without_vision_pictures_are_recorded_unread(self):
        from test_textdocs import extract
        _, coverage, _, model = extract("process.docx", build(), vision=False)
        self.assertEqual({r["status"] for r in coverage if r["task"].startswith("picture")}, {"skipped"})
        self.assertFalse(any(q.region != "text" for q in model.queries))

if __name__ == "__main__":
    unittest.main()
