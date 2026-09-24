import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils.cell import range_boundaries
except ImportError:
    openpyxl = None

SCRIPT = Path(__file__).resolve().parent.parent / "scripts" / "make_synthetic_samples.py"

def load_generator():
    spec = importlib.util.spec_from_file_location("make_synthetic_samples", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

@unittest.skipUnless(openpyxl, "openpyxl not installed (pip install -e .[dev])")
class SyntheticSampleTests(unittest.TestCase):
    """The answer keys must describe the generated files exactly."""

    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory()
        cls.dest = Path(cls.directory.name)
        load_generator().main(["--dest", str(cls.dest)])

    @classmethod
    def tearDownClass(cls):
        cls.directory.cleanup()

    def test_workbook_regions_cover_every_cell_exactly_once(self):
        key = json.loads((self.dest / "messy-workbook.expected.json").read_text())
        wb = openpyxl.load_workbook(self.dest / "messy-workbook.xlsx")
        self.assertEqual(set(key["sheets"]), set(wb.sheetnames))
        for name, sheet in key["sheets"].items():
            ws = wb[name]
            self.assertEqual(ws.sheet_state == "hidden", sheet["hidden"], name)
            boxes = [range_boundaries(r["range"] if ":" in r["range"] else f"{r['range']}:{r['range']}")
                     for r in sheet["regions"]]
            filled = [(c.column, c.row) for row in ws.iter_rows() for c in row if c.value is not None]
            for column, row in filled:
                owners = [b for b in boxes if b[0] <= column <= b[2] and b[1] <= row <= b[3]]
                self.assertEqual(len(owners), 1, f"{name}!{openpyxl.utils.get_column_letter(column)}{row}")
            for region, (x0, y0, x1, y1) in zip(sheet["regions"], boxes):
                self.assertTrue(any(x0 <= c <= x1 and y0 <= r <= y1 for c, r in filled), f"{name} {region['range']} is empty")

    def test_formula_has_no_cached_value(self):
        wb = openpyxl.load_workbook(self.dest / "messy-workbook.xlsx", data_only=True)
        self.assertIsNone(wb["Summary"]["C8"].value)
        self.assertEqual(openpyxl.load_workbook(self.dest / "messy-workbook.xlsx")["Summary"]["C8"].value, "=C6+C7")

    def test_csv_regions_cover_every_nonblank_line(self):
        key = json.loads((self.dest / "messy-export.expected.json").read_text())
        lines = (self.dest / "messy-export.csv").read_text().splitlines()
        regions = key["sheets"]["(csv)"]["regions"]
        for number, line in enumerate(lines, 1):
            owners = [r for r in regions if r["rows"][0] <= number <= r["rows"][1]]
            self.assertEqual(len(owners), 1 if line.strip() else 0, f"line {number}: {line!r}")

    def test_output_is_deterministic(self):
        with tempfile.TemporaryDirectory() as other:
            load_generator().main(["--dest", other])
            for name in ["messy-workbook.xlsx", "messy-workbook.expected.json", "messy-export.csv", "messy-export.expected.json"]:
                self.assertEqual((self.dest / name).read_bytes(), (Path(other) / name).read_bytes(), name)

if __name__ == "__main__":
    unittest.main()
