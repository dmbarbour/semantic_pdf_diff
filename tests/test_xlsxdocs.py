"""Excel workbooks (xlsxdocs.py; the adapters plan, milestone 5, "Excel, in detail"): sheets as pages, regions found,
cells as Excel shows them, tables on the grid, claims placed by sheet and cells."""
import stubs  # noqa: F401 (a clean environment)
import io
import unittest

try:
    import docx  # noqa: F401 (the office extra)
    import openpyxl
except ImportError:
    openpyxl = None

def workbook():
    """A workbook: a sheet with a merged title, a note, a table with a two-row header (one cell merged across), a
    formula left uncalculated and a comment; a sheet of label/value pairs over a table, two tables pressed together
    and a defined table too long to read row by row; and a hidden sheet."""
    from openpyxl.comments import Comment
    from openpyxl.worksheet.table import Table
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Pump Station PS-3 Design Summary"
    ws.merge_cells("A1:E1")
    ws["A2"] = "Values marked TBC are provisional."
    for r, row in enumerate([["Parameter", "Measured", None, "Design"], [None, "value", "±", None],
                             ["Flow (L/s)", 118, 3, 120], ["Head (m)", 24.1, 0.4, 25.5]], 4):
        for c, value in enumerate(row, 1):
            if value is not None:
                ws.cell(r, c, value)
    ws.merge_cells("B4:C4")
    ws["A8"], ws["B8"] = "Total head (m)", "=B6+B7"  # never calculated: openpyxl stores no cached value
    ws["B7"].number_format = "0.00"
    ws["A6"].comment = Comment("Retested in June.", "Ana")
    two = wb.create_sheet("Polar")
    for r, row in enumerate([["Configuration", "default"], ["Reynolds", 8100000],
                             ["alpha [deg]", "c_l", "c_d"], [-10, -0.82, 0.0123], [0, 0.31, 0.0071]], 1):
        for c, value in enumerate(row, 1):
            two.cell(r, c, value)
    for r, row in enumerate([["Option", "Capex (k$)", "Opex (k$/yr)"], ["A", 410, 22], ["B", 380, 30, "Year", "MWh"],
                             [None, None, None, 2025, 410]], 8):  # a second table pressed against the first
        for c, value in enumerate(row, 1):
            if value is not None:
                two.cell(r, c, value)
    rows = [["ID", "Requirement", "Value"]] + [[f"R-{k:03d}", f"Requirement {k}", k * 3] for k in range(1, 61)]
    for r, row in enumerate(rows, 14):
        for c, value in enumerate(row, 1):
            two.cell(r, c, value)
    two.add_table(Table(displayName="Requirements", ref=f"A14:C{13 + len(rows)}"))
    old = wb.create_sheet("Rev B (superseded)")
    old["A1"], old["B1"] = "Design flow (L/s)", 110
    old["A2"], old["B2"] = "Static head (m)", 18
    old.sheet_state = "hidden"
    raw = io.BytesIO()
    wb.save(raw)
    return raw.getvalue()

@unittest.skipIf(openpyxl is None, "python-docx and openpyxl aren't installed (the office extra)")
class Reading(unittest.TestCase):
    def test_sheets_regions_and_cells_as_excel_shows_them(self):
        from semantic_pdf_diff.xlsxdocs import read_xlsx
        doc = read_xlsx(workbook())
        self.assertEqual([title for _, _, _, title in doc.headings], ["Summary", "Polar", "Rev B (superseded)"])
        found = [(g.sheet, g.ref, g.kind) for g in doc.regions]
        self.assertEqual(found[:4], [("Summary", "A1", "text"), ("Summary", "A2", "text"), ("Summary", "A4:D8", "table"),
                                     ("Polar", "A1:B2", "pairs")])
        summary = next(b for b in doc.blocks if b.kind == "table")
        self.assertEqual(summary.rows[0], ["Parameter", "Measured > value", "Measured > ±", "Design"])
        self.assertEqual(summary.rows[1], ["Flow (L/s)", "118", "3", "120"])
        self.assertEqual(summary.rows[2], ["Head (m)", "24.10", "0.4", "25.5"])  # its number format: two decimals
        self.assertEqual(summary.rows[3][:2], ["Total head (m)", "[formula =B6+B7, not calculated]"])
        texts = [t for _, t in doc.lines]
        self.assertIn("Comment by Ana on A6: Retested in June.", texts)
        self.assertIn("Reynolds: 8100000", texts)  # a label and its value over the table
        self.assertIn("(Hidden sheet)", texts)
        polar = [b for b in doc.blocks if b.kind == "table"][1]
        self.assertEqual(polar.rows[0], ["alpha [deg]", "c_l", "c_d"])
        self.assertEqual(polar.grid.title, "Configuration: default; Reynolds: 8100000")  # the pairs heading it

    def test_an_ambiguous_region_is_recorded_and_a_long_table_kept_whole(self):
        from semantic_pdf_diff.xlsxdocs import read_xlsx
        doc = read_xlsx(workbook())
        unread = [what for *_, what in doc.images]
        self.assertTrue(any(w.startswith("skipped: ambiguous sheet layout (Polar!A8:E11)") for w in unread))
        long = next(b for b in doc.blocks if b.kind == "table" and b.grid.place == "Polar!A14:C74")
        self.assertEqual((len(long.rows), long.grid.place, long.grid.columns), (61, "Polar!A14:C74", ["A", "B", "C"]))
        self.assertEqual((long.grid.names[0], long.grid.rows[0], long.grid.labels), ("15", ["R-001", "Requirement 1", "3"],
                                                                                      ["ID", "Requirement", "Value"]))

    def test_a_workbook_is_extracted_with_cell_locators(self):
        import sys
        from pathlib import Path
        sys.path.insert(0, str(Path(__file__).parent))
        from test_textdocs import extract
        from semantic_pdf_diff.models import XlsxLocator
        evidence, coverage, sections, _ = extract("book.xlsx", workbook())
        values = {e.value: e for e in evidence}
        located = values["118"].locator
        self.assertIsInstance(located, XlsxLocator)
        self.assertEqual((located.page, located.sheet, located.cells, located.region), (1, "Summary", "A6:D6", "table"))
        self.assertEqual(values["110"].locator.sheet, "Rev B (superseded)")
        self.assertEqual([s.heading_path for s in sections], [["Summary"], ["Polar"], ["Rev B (superseded)"]])

CSV = ("# Exported by SCADA historian\n"
       "Site: PS-3,Interval: daily\n"
       "\n"
       "Date,Volume pumped (m3),Energy (kWh)\n"
       "2026-07-01,8420,1712\n"
       "2026-07-02,8390,\"1,705\"\n"
       "\n"
       "Alarm,Count,Note\n"
       "Pump trip,1,\"tripped twice,\nreset by hand\"\n"
       "Comms loss,5,\n")

class Csv(unittest.TestCase):
    """CSV files (xlsxdocs.read_csv): a one-sheet workbook of text cells, placed by the file's lines and fields; no
    extra needed."""
    def test_preamble_and_tables_placed_by_lines_and_fields(self):
        from semantic_pdf_diff.xlsxdocs import read_csv
        doc = read_csv(CSV.encode())
        self.assertEqual([(g.ref, g.kind) for g in doc.regions], [("A1", "text"), ("A2:B2", "text"), ("A4:C6", "table"),
                                                                  ("A8:C10", "table")])
        texts = [t for _, t in doc.lines]
        self.assertEqual(texts[:3], ["# Exported by SCADA historian", "Site: PS-3", "Interval: daily"])
        alarms = [b for b in doc.blocks if b.kind == "table"][1]
        self.assertEqual(alarms.rows[1], ["Pump trip", "1", "tripped twice, reset by hand"])
        self.assertEqual(doc.places[alarms.row_lines[1]], ((9, 10), (1, 3)))  # the quoted field spans two lines
        self.assertEqual(doc.places[alarms.row_lines[2]], ((11, 11), (1, 3)))
        daily = [b for b in doc.blocks if b.kind == "table"][0]
        self.assertEqual((daily.grid.place, daily.grid.rows[1]), ("A4:C6", ["2026-07-02", "8390", "1,705"]))

    def test_semicolons_and_windows_1252(self):
        from semantic_pdf_diff.xlsxdocs import read_csv
        doc = read_csv("Pumpe;Förderhöhe (m);Leistung (kW)\nP-1;24,1;22\nP-2;25,5;22\n".encode("cp1252"))
        table, = [b for b in doc.blocks if b.kind == "table"]
        self.assertEqual(table.rows, [["Pumpe", "Förderhöhe (m)", "Leistung (kW)"], ["P-1", "24,1", "22"],
                                      ["P-2", "25,5", "22"]])

    def test_a_csv_file_is_extracted_with_line_and_field_locators(self):
        import sys
        from pathlib import Path
        sys.path.insert(0, str(Path(__file__).parent))
        from test_textdocs import extract
        from semantic_pdf_diff.models import CsvLocator
        evidence, coverage, sections, _ = extract("export.csv", CSV.encode())
        located = {e.value: e.locator for e in evidence}["8420"]
        self.assertIsInstance(located, CsvLocator)
        self.assertEqual((located.file_lines, located.fields, located.region), ((5, 5), (1, 3), "table"))
        self.assertEqual(evidence[0].derivation[0].step[:4], "csv-")

class Readers(unittest.TestCase):
    def test_every_format_read_has_a_reader_version(self):
        from semantic_pdf_diff.extract import IMAGE_EXTENSIONS, READERS, TEXT_EXTENSIONS
        self.assertEqual(set(TEXT_EXTENSIONS) | set(IMAGE_EXTENSIONS) | {".pdf"}, set(READERS))

class RulesModel:
    """The text tests' model, answering a table's rules query with the given rules (keyed by the table's place), or
    with each row read by itself."""
    def __init__(self, answers, reviews=None, **settings):
        import sys
        from pathlib import Path
        sys.path.insert(0, str(Path(__file__).parent))
        from test_textdocs import Model
        self.base, self.answers, self.asked, self.reviews, self.reviewed = Model(**settings), answers, [], reviews or {}, []
        self.s = self.base.s

    def __getattr__(self, name):
        return getattr(self.base, name)

    def ask(self, prompt, schema, images=(), key=None):
        from semantic_pdf_diff.tablerules import Review, Rules
        if schema is Review:
            self.reviewed.append(prompt)
            place = prompt.split("TABLE: ", 1)[1].split(";", 1)[0].split("\n", 1)[0]
            said = self.reviews.get(place, {"verdict": "keep"})  # a list: one answer a review, then keep
            if isinstance(said, list):
                said = said.pop(0) if said else {"verdict": "keep"}
            return Review.model_validate(said)
        if schema is not Rules:
            return self.base.ask(prompt, schema, images, key)
        self.asked.append(prompt)
        place = prompt.split("TABLE: ", 1)[1].split(";", 1)[0].split("\n", 1)[0]
        return Rules.model_validate(self.answers.get(place, {"reading": "rows"}))

REQUIREMENTS = {"binding": "values of requirements, by ID", "reading": "rules", "subject": "requirements",
                "claims": [{"entity": "{A}", "attribute": "{C.header}", "value": "{C}", "number": True}],
                "examples": [{"row": "15", "claims": [{"entity": "R-001", "attribute": "Value", "value": "3"}]}]}

@unittest.skipIf(openpyxl is None, "python-docx and openpyxl aren't installed (the office extra)")
class Rules(unittest.TestCase):
    def extract(self, answers, reviews=None, **settings):
        import tempfile
        from pathlib import Path
        from semantic_pdf_diff.extract import Job, reader_for, run_jobs
        model = RulesModel(answers, reviews, **settings)
        data = workbook()
        with tempfile.TemporaryDirectory() as d:
            job = Job("sha256:" + "a" * 64 + ".xlsx", lambda: data, reader=reader_for(".xlsx"))
            run_jobs([[job]], Path(d), model)
        evidence, coverage = job.state["result"]
        return evidence, coverage, model

    def test_every_table_is_asked_how_its_read(self):
        evidence, coverage, model = self.extract({"Polar!A14:C74": REQUIREMENTS})
        self.assertEqual([q.split("TABLE: ", 1)[1].split(";")[0].split("\n")[0] for q in model.asked],
                         ["Summary!A4:D8", "Polar!A3:C5", "Polar!A14:C74", "Rev B (superseded)!A1:B2"])
        long = model.asked[2]
        self.assertIn("SIZE: 60 rows, 3 columns", long)
        self.assertIn("CONTEXT (for reference only: do not extract claims from it):\nAbove the table: ...", long)
        self.assertIn("Within: Polar", long)  # the context a row of it would be given
        ruled = [e for e in evidence if e.derivation[-1].step == "table-rules"]
        self.assertEqual(len(ruled), 60)
        first = next(e for e in ruled if e.entity == "R-001")
        self.assertEqual((first.attribute, first.value, first.quote, first.quote_verified), ("Value", "3", "R-001 | 3", True))
        self.assertEqual((first.locator.sheet, first.locator.cells), ("Polar", "A15:C15"))
        row = next(r for r in coverage if r["task"] == "rules:p2:1")
        self.assertEqual((row["status"], row["claims"]), ("complete", 60))
        self.assertIn("[binding: values of requirements, by ID]", row["issues"][0])
        self.assertIn("How a claim is bound", long)
        self.assertEqual(sum(r["task"].startswith("table:p1:0:") for r in coverage), 3)  # "rows": each by itself
        self.assertFalse([r for r in coverage if r["task"].startswith("table:p2:1:")])  # no row asked by itself

    def test_rules_that_miss_their_examples_are_asked_again_then_rows_read(self):
        wrong = {**REQUIREMENTS, "examples": [{"row": "15", "claims": [{"value": "4"}]}]}
        evidence, coverage, model = self.extract({"Polar!A14:C74": wrong})
        self.assertEqual(len(model.asked), 5)
        self.assertIn("ITS PROBLEMS:\n- row 15: the templates give 3, not 4", model.asked[3])
        tasks = [r["task"] for r in coverage]
        self.assertIn("rules:p2:1:again", tasks)
        self.assertEqual(sum(t.startswith("table:p2:1:") for t in tasks), 60)  # every row read by itself

    def test_a_review_shows_the_outcome_and_may_revise_the_rules(self):
        from semantic_pdf_diff import tablerules
        revised = {**REQUIREMENTS, "claims": [{"entity": "{A}", "attribute": "required value", "value": "{C}",
                                               "number": True}]}
        tablerules.REVIEW = 10
        try:
            evidence, coverage, model = self.extract({"Polar!A14:C74": REQUIREMENTS}, {"Polar!A14:C74": [{
                "verdict": "revise", "problems": ["the attribute is the column's name"], "rules": revised}]})
            broken = {**revised, "claims": [{"entity": "{A}", "attribute": "x", "value": "{Q}"}]}
            _, kept, _ = self.extract({"Polar!A14:C74": REQUIREMENTS}, {"Polar!A14:C74": {"verdict": "revise",
                                                                                         "rules": broken}})
            tablerules.REVIEW = 3  # a reviewer that never keeps: stopped at the cap, its last revision used
            capped, endless, _ = self.extract({"Polar!A14:C74": REQUIREMENTS}, {"Polar!A14:C74": [
                {"verdict": "revise", "problems": ["again"], "rules": {**revised, "claims": [
                    {**revised["claims"][0], "attribute": f"value {k}"}]}} for k in range(5)]})
        finally:
            tablerules.REVIEW = None
        review = next(q for q in model.reviewed if "TABLE: Polar!A14:C74" in q)
        self.assertIn("WHAT THEY GAVE:\nrow 15: R-001 | Requirement 1 | 3\n  R-001 | Value | 3", review)
        ruled = [e for e in evidence if e.derivation[-1].step == "table-rules"]
        self.assertEqual({e.attribute for e in ruled}, {"required value"})
        self.assertEqual(ruled[0].derivation[-2].detail, "revised 1 time")
        row = next(r for r in coverage if r["task"] == "rules:p2:1")
        self.assertIn("Reviewed: revised (the attribute is the column's name)", row["issues"])
        self.assertIn("Reviewed: kept", row["issues"])  # shown again after the revision, and kept
        self.assertEqual(sum("TABLE: Polar!A14:C74" in q for q in model.reviewed), 2)
        tablerules.REVIEW = 10
        try:  # a revision giving back the same rules stops the review at once
            _, still, model = self.extract({"Polar!A14:C74": REQUIREMENTS}, {"Polar!A14:C74": {
                "verdict": "revise", "problems": ["wants what the templates can't say"], "rules": REQUIREMENTS}})
        finally:
            tablerules.REVIEW = None
        self.assertEqual(sum("TABLE: Polar!A14:C74" in q for q in model.reviewed), 1)
        self.assertTrue(any(i.startswith("Reviewed: a revision changing nothing") for i in
                            next(r for r in still if r["task"] == "rules:p2:1")["issues"]))
        last = next(e for e in capped if e.derivation[-1].step == "table-rules")
        self.assertEqual(last.derivation[-2].detail, "revised 3 times, the cap reached")
        self.assertIn("Review cap (3) reached: the last revision used",
                      next(r for r in endless if r["task"] == "rules:p2:1")["issues"])
        row = next(r for r in kept if r["task"] == "rules:p2:1")
        self.assertTrue(any(i.startswith("Reviewed: a revision with problems (no column Q") for i in row["issues"]))

    def test_a_late_answer_reads_its_own_sheets_rows(self):
        # Sheet 1's rules answer ("rows") comes only after sheet 2's rows have been asked; its rows must still be read
        # as sheet 1's (code review 2026-10-08, B1: they were filed under the sheet the loop had reached).
        import tempfile
        import threading
        from pathlib import Path
        from types import SimpleNamespace
        from semantic_pdf_diff.extract import Job, reader_for, run_jobs
        from semantic_pdf_diff.models import Extraction, Settings
        from semantic_pdf_diff.tablerules import Rules
        wb = openpyxl.Workbook()
        for title, header, rows in [("Pumps", ["Tag", "Flow (L/s)"], [["P-1", 120], ["P-2", 95]]),
                                    ("Fans", ["Fan", "Speed (rpm)"], [["F-1", 900], ["F-2", 1200], ["F-3", 1500]])]:
            ws = wb.active if title == "Pumps" else wb.create_sheet(title)
            ws.title = title
            for r, row in enumerate([header] + rows, 1):
                for c, value in enumerate(row, 1):
                    ws.cell(r, c, value)
        out = io.BytesIO()
        wb.save(out)
        sheet2_read = threading.Event()

        class Client:  # answers on worker threads (the dispatcher's prepare/send split), so they can come late
            s = Settings(vision=False, concurrency=4)
            calls, cache_hits, usage = 0, 0, {}
            asked, late = [], None

            def prepare(self, prompt, schema, images, key):
                return SimpleNamespace(prompt=prompt, schema=schema, key=key, query=None)

            def cached(self, request):
                return None

            def save(self, request, value):
                pass

            def send(self, request):
                return self.ask(request.prompt, request.schema, key=request.key)

            def ask(self, prompt, schema, images=(), key=None):
                self.asked.append((key[3], prompt))
                if schema is Rules:
                    if "TABLE: Pumps!" in prompt:
                        self.late = sheet2_read.wait(5)
                    return Rules(reading="rows")
                if schema is Extraction and key[3].startswith("table:p2:"):
                    sheet2_read.set()
                if schema is Extraction:
                    return Extraction(claims=[], complete=True)
                return schema.model_validate({})

        client = Client()
        with tempfile.TemporaryDirectory() as d:
            job = Job("sha256:" + "b" * 64 + ".xlsx", lambda: out.getvalue(), reader=reader_for(".xlsx"))
            run_jobs([[job]], Path(d), client)
        _, coverage = job.state["result"]
        self.assertTrue(client.late)  # the answer did come after sheet 2's rows were asked
        rows = {task: prompt for task, prompt in client.asked if task.startswith("table:")}
        self.assertEqual(sorted(t for t in rows if t.startswith("table:p1:")), ["table:p1:0:0", "table:p1:0:1"])
        self.assertTrue(all("P-" in rows[t] for t in rows if t.startswith("table:p1:")))
        self.assertFalse(any("P-" in rows[t] for t in rows if t.startswith("table:p2:")))
        self.assertEqual(len([r for r in coverage if r["task"].startswith("table:")]), 5)  # no tag shared

    def test_tables_asked_from_a_size_or_none(self):
        _, coverage, model = self.extract({"Polar!A14:C74": REQUIREMENTS}, table_rules=50)
        self.assertEqual(len(model.asked), 1)  # only the table of more than 50 rows
        _, coverage, model = self.extract({}, table_rules=None)
        self.assertEqual(model.asked, [])
        self.assertEqual(sum(r["task"].startswith("table:p2:1:") for r in coverage), 60)

if __name__ == "__main__":
    unittest.main()
