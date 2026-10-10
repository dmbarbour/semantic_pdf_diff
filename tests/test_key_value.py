"""Blocks of two columns, proposed as key-value lists and confirmed by the model (code review 2026-10-08, B5): a block
standing alone ("Pump | P-2", "Flow | 95 L/s", "Head | 30 m") was a table headed by its first pair, so "P-2" was
never read as a value and the other rows were read under the labels "Pump" and "P-2"."""
import stubs  # noqa: F401 (a clean environment)
import contextlib
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from semantic_pdf_diff import keyvalue, tablerules
from semantic_pdf_diff.llm import CallLimitReached, ModelFailure
from test_textdocs import Model

PAIRS = [["Pump", "P-2"], ["Flow", "95 L/s"], ["Head", "30 m"]]

class Answering(Model):
    """The text tests' model, answering the key-value check as told (a reading, or an error raised)."""
    def __init__(self, check, **settings):
        super().__init__(**settings)
        self.check, self.checked = check, []

    def ask(self, prompt, schema, images=(), key=None):
        if schema is keyvalue.Check:
            self.checked.append((prompt, list(images)))
            if isinstance(self.check, Exception):
                raise self.check
            return keyvalue.Check(reading=self.check, why="as a form")
        return super().ask(prompt, schema, images, key)

def run(name, data, check, **settings):
    """(evidence, coverage, the model) of one file read by its reader, the check answered as `check`."""
    from semantic_pdf_diff.jobs import Job, reader_for, run_jobs
    model = Answering(check, **settings)
    with tempfile.TemporaryDirectory() as d:
        job = Job("sha256:" + "c" * 64 + Path(name).suffix, lambda: data, reader=reader_for(Path(name).suffix))
        run_jobs([[job]], Path(d), model)
    return *job.state["result"], model

def workbook(rows, defined=False):
    import openpyxl
    from openpyxl.worksheet.table import Table
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Pump"
    for row in rows:
        sheet.append(row)
    if defined:
        sheet.add_table(Table(displayName="Pump", ref=f"A1:B{len(rows)}"))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()

def check_row(coverage):
    return next(r for r in coverage if r["task"].startswith("key-value:"))

def office():
    try:
        import docx  # noqa: F401
        import openpyxl  # noqa: F401
    except ImportError:
        return False
    return True

class Candidates(unittest.TestCase):
    def grid(self, rows, columns=2):
        return tablerules.grid([tablerules.letter(k) for k in range(1, columns + 1)], rows[0], rows[1:],
                               [str(k) for k in range(2, len(rows) + 1)], list(range(2, len(rows) + 1)))

    def test_two_columns_with_no_numbers_on_the_left_are_proposed(self):
        self.assertTrue(keyvalue.candidate(self.grid(PAIRS)))
        self.assertTrue(keyvalue.candidate(self.grid([["Tag", "Flow (L/s)"], ["P-1", "120"]])))  # the model decides
        self.assertFalse(keyvalue.candidate(self.grid(PAIRS), marked=True))
        self.assertFalse(keyvalue.candidate(self.grid([["Year", "Output"], ["2020", "5"], ["2021", "6"]])))
        self.assertFalse(keyvalue.candidate(self.grid([["2024-01-01", "start"], ["Flow", "95"]])))
        self.assertFalse(keyvalue.candidate(self.grid([r + ["x"] for r in PAIRS], columns=3)))
        self.assertFalse(keyvalue.candidate(self.grid([["Pumps", "Pumps"], ["Flow", "95"]])))  # a title over both
        self.assertFalse(keyvalue.candidate(self.grid([["Pump", ""], ["Flow", "95"]])))
        self.assertFalse(keyvalue.candidate(self.grid(PAIRS[:1])))  # no row below
        self.assertFalse(keyvalue.candidate(self.grid([["Rating", "(psi)"], ["Rating", "(psi)"]])))  # a PDF's one row
        self.assertFalse(keyvalue.candidate(None))

    def test_the_check_states_the_proposal_and_shows_every_row_numbered(self):
        asked = keyvalue.question(self.grid(PAIRS), "Above the table: ...the pump's data", image=True)
        self.assertIn('Our proposal, a guess from its layout (two columns, no numbers in column A): a key-value list',
                      asked)
        self.assertIn("1 | Pump | P-2\n2 | Flow | 95 L/s\n3 | Head | 30 m", asked)
        self.assertIn("Above the table: ...the pump's data", asked)
        self.assertTrue(asked.endswith(keyvalue.IMAGED))
        self.assertEqual(keyvalue.pairs(self.grid(PAIRS)), ["Pump: P-2", "Flow: 95 L/s", "Head: 30 m"])

@unittest.skipUnless(office(), "python-docx and openpyxl aren't installed (the office extra)")
class Sheets(unittest.TestCase):
    def test_a_key_value_list_is_read_as_one(self):
        evidence, coverage, model = run("p.xlsx", workbook(PAIRS), "key-value")
        prompt, images = model.checked[0]
        self.assertIn("1 | Pump | P-2", prompt)
        self.assertEqual(images, [])  # a sheet has no image: its text alone
        self.assertEqual(check_row(coverage)["status"], "complete")
        self.assertIn("a key-value list, the model's reading (as a form)", check_row(coverage)["issues"][0])
        read = [q for q in model.queries if q.region == "table" or "Pump: P-2" in q.data]
        self.assertEqual([q.region for q in read], ["text"])  # one text task, no table rows
        self.assertIn("Pump: P-2\n\nFlow: 95 L/s\n\nHead: 30 m", read[0].data)
        claim = next(e for e in evidence if e.value == "95")
        self.assertEqual(claim.locator.cells, "A2:B2")  # placed on its own row
        self.assertEqual([s.step for s in claim.derivation], ["xlsx-table", "key-value-check", "model-extraction"])

    def test_a_table_the_model_reads_as_one_is_read_as_before_its_claims_traced(self):
        evidence, coverage, model = run("p.xlsx", workbook(PAIRS), "table")
        rows = [q.data for q in model.queries if q.region == "table"]
        self.assertEqual(len(rows), 2)
        self.assertIn('Header: ["Pump", "P-2"]', rows[0])
        self.assertIn("a table, the model's reading", check_row(coverage)["issues"][0])
        claim = next(e for e in evidence if e.value == "95")
        self.assertEqual([s.step for s in claim.derivation], ["xlsx-table", "key-value-check", "model-extraction"])

    def test_a_failed_check_reads_our_proposal(self):
        evidence, coverage, model = run("p.xlsx", workbook(PAIRS), ModelFailure("timed out"))
        self.assertEqual(check_row(coverage)["status"], "partial")
        self.assertIn("our proposal, not confirmed (timed out)", check_row(coverage)["issues"][0])
        self.assertFalse([q for q in model.queries if q.region == "table"])
        self.assertTrue(any("Pump: P-2" in q.data for q in model.queries))
        claim = next(e for e in evidence if e.value == "95")
        self.assertIn("not confirmed", claim.derivation[1].detail)
        _, coverage, _ = run("p.xlsx", workbook(PAIRS), "neither")  # an answer naming neither: as a failure
        self.assertEqual(check_row(coverage)["status"], "partial")

    def test_an_unreached_check_reads_nothing_until_the_next_run(self):
        _, coverage, model = run("p.xlsx", workbook(PAIRS), CallLimitReached("the call limit"))
        self.assertEqual(check_row(coverage)["status"], "not_reached")
        self.assertFalse([q for q in model.queries if "P-2" in q.data])

    def test_a_defined_table_and_a_table_of_numbers_arent_asked(self):
        _, _, model = run("p.xlsx", workbook(PAIRS, defined=True), "key-value")
        self.assertEqual(model.checked, [])
        _, _, model = run("p.xlsx", workbook([["Year", "Output (MWh)"], [2020, 5], [2021, 6]]), "key-value")
        self.assertEqual(model.checked, [])

    def test_a_csv_files_block_is_asked(self):
        _, _, model = run("p.csv", b"Pump,P-2\nFlow,95 L/s\nHead,30 m\n", "key-value")
        self.assertEqual(len(model.checked), 1)
        self.assertTrue(any("Flow: 95 L/s" in q.data for q in model.queries))

@unittest.skipUnless(office(), "python-docx and openpyxl aren't installed (the office extra)")
class Word(unittest.TestCase):
    def document(self, header_marked=False):
        import docx
        from docx.oxml import OxmlElement
        d = docx.Document()
        d.add_paragraph("Pump data")
        table = d.add_table(rows=len(PAIRS), cols=2)
        for r, (key, value) in enumerate(PAIRS):
            table.cell(r, 0).text, table.cell(r, 1).text = key, value
        if header_marked:
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        out = io.BytesIO()
        d.save(out)
        return out.getvalue()

    def test_a_two_column_table_is_asked_unless_its_header_is_marked(self):
        evidence, _, model = run("p.docx", self.document(), "key-value")
        self.assertEqual(len(model.checked), 1)
        self.assertIn("Above the table: ...Pump data", model.checked[0][0])
        pairs, = [q for q in model.queries if "Pump: P-2" in q.data]
        self.assertIn("Above the table: ...Pump data", pairs.context)  # as a table's row: its caption, the text above
        claim = next(e for e in evidence if e.value == "95")
        self.assertEqual([s.step for s in claim.derivation], ["docx-table", "key-value-check", "model-extraction"])
        _, _, model = run("p.docx", self.document(header_marked=True), "key-value")
        self.assertEqual(model.checked, [])

class Pdf(unittest.TestCase):
    """The PDF job's table path, on a detected table of two columns (find_tables patched)."""
    BOXES = tuple((20, 40 + 12 * k, 280, 52 + 12 * k) for k in range(len(PAIRS)))

    def extract(self, check):
        import pymupdf
        from semantic_pdf_diff.tasks import ExtractQuery
        from semantic_pdf_diff.jobs import extract_pdf
        from semantic_pdf_diff.schema import Extraction
        from semantic_pdf_diff.settings import Settings
        from semantic_pdf_diff.provenance import content_id
        from stubs import situating_answer

        class Table:
            bbox, cells = (20, 40, 280, 40 + 12 * len(PAIRS)), []
            def extract(self): return [list(r) for r in PAIRS]

        class Client:
            s = Settings(vision=False, table_structure=False, table_rules=None, refinement_depth=0)
            calls, cache_hits, usage = 0, 0, {}

            def __init__(self):
                self.checked, self.asked = [], {}

            def ask(self, prompt, schema, images=(), key=None):
                if situating_answer(prompt):
                    return schema.model_validate(situating_answer(prompt))
                if schema is keyvalue.Check:
                    self.checked.append(list(images))
                    return keyvalue.Check(reading=check)
                query = ExtractQuery.read(prompt)
                self.asked[key[3]] = query.data
                claims = [{"entity": "pump", "attribute": "flow", "value": "95", "unit": "L/s", "kind": "text",
                           "quote": "95 L/s", "confidence": 0.9}] if "95 L/s" in query.data else []
                return Extraction(claims=claims, complete=True)

        find = lambda page, *a, **k: type("T", (), {"tables": [Table()]})()
        with contextlib.ExitStack() as stack, tempfile.TemporaryDirectory() as d:
            stack.enter_context(patch.object(pymupdf.Page, "find_tables", find))
            stack.enter_context(patch("semantic_pdf_diff.extract.Marks", lambda page: None))
            stack.enter_context(patch("semantic_pdf_diff.extract.row_boxes", lambda table, rows: self.BOXES))
            stack.enter_context(contextlib.redirect_stderr(io.StringIO()))
            doc = pymupdf.open()
            doc.new_page(width=300, height=300)
            path = Path(d) / "t.pdf"
            doc.save(path)
            client = Client()
            evidence, coverage = extract_pdf(path, content_id(path.read_bytes(), path.name), Path(d), client)
        return client, evidence, coverage

    def test_a_key_value_list_is_asked_with_its_crop_and_read_by_line(self):
        client, evidence, coverage = self.extract("key-value")
        self.assertEqual(len(client.checked[0]), 1)  # the table's crop
        self.assertEqual(check_row(coverage)["task"], "key-value:p1:0")
        self.assertEqual(client.asked["text:p1:kv0"], "Pump: P-2\n\nFlow: 95 L/s\n\nHead: 30 m")
        self.assertFalse([t for t in client.asked if t.startswith("table:")])
        claim = next(e for e in evidence if e.value == "95")
        self.assertEqual(tuple(claim.locator.bbox), self.BOXES[1])  # its own row's box
        self.assertEqual([s.step for s in claim.derivation],
                         ["pdf-table-detection", "key-value-check", "model-extraction"])

    def test_a_table_is_read_as_before(self):
        client, evidence, _ = self.extract("table")
        self.assertEqual(sorted(t for t in client.asked if t.startswith("table:")), ["table:p1:0:0", "table:p1:0:1"])
        self.assertIn('Header: ["Pump", "P-2"]', client.asked["table:p1:0:0"])
        claim = next(e for e in evidence if e.value == "95")
        self.assertEqual([s.step for s in claim.derivation],
                         ["pdf-table-detection", "key-value-check", "model-extraction"])

class Refined(unittest.TestCase):
    def test_a_refined_parts_context_is_its_own_unless_given(self):
        # a key-value list's text task is given a table's context; its parts keep it. A text block's parts compute
        # their own, as before (B5's first version passed the parent's on, changing every refined text request).
        from types import SimpleNamespace
        from semantic_pdf_diff.settings import Settings
        from semantic_pdf_diff.tasks import TaskCore
        job = SimpleNamespace(content="sha256:" + "d" * 64 + ".txt", on_task=None, state={"pending": 0})
        client = SimpleNamespace(s=Settings(refinement_depth=1))
        core = TaskCore(job, Path("."), client, None, None, "t", None, {})
        core.reader = SimpleNamespace(for_text=lambda page, segments, text: f"context of {segments[0][1]}")
        sent = []

        def consume(page, box, task, text, context="", then=None, **_):
            sent.append((task, context))
            if then and ":r" not in task:
                then("partial")
        core.consume = consume
        segments = [((0, 1, 1, 2), "first"), ((0, 2, 1, 3), "second")]
        core.text_task(1, segments, "text:p1:0")
        self.assertEqual(sent, [("text:p1:0", "context of first"), ("text:p1:0:r0", "context of first"),
                                ("text:p1:0:r1", "context of second")])
        sent.clear()
        core.text_task(1, segments, "text:p1:kv0", context="a table's")
        self.assertEqual([c for _, c in sent], ["a table's"] * 3)

if __name__ == "__main__":
    unittest.main()
